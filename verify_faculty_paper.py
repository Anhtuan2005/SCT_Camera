import hashlib
import re
import sys
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn
from lxml import etree


TEMPLATE = Path(r"E:\SCT_Camera\.codex_work\faculty_template\faculty_template.docx")
SOURCE = Path(r"E:\Paper_NCKH.docx")
FINAL = Path(r"E:\SCT_Camera\.codex_work\faculty_template\paper_faculty_v3.docx")
PDF = Path(r"E:\SCT_Camera\.codex_work\faculty_template\render_v3\paper_faculty_v3.pdf")
EXPECTED_TEMPLATE_SHA = "a55600f881167f2fa37cb23995d8957102aece4393788cb8d2c9dcd781193d97"

REQUIRED_STYLES = [
    "TenBaiBao", "TacGia", "TomTat", "TuKhoa", "MucChinh", "MucPhu1",
    "MucPhu2", "NoiDung", "TenBang", "TenHinh", "TieuDeBang",
    "NoiDungBang", "TaiLieuThamKhao",
]


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def related_image_hashes(doc):
    hashes = []
    for p in doc.paragraphs:
        for blip in p._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            hashes.append(sha256(doc.part.related_parts[rid].blob))
    return hashes


def node_signature(node):
    return (
        etree.QName(node).localname,
        (node.text or "").strip(),
        tuple(sorted((etree.QName(k).localname, v) for k, v in node.attrib.items())),
        tuple(node_signature(child) for child in node),
    )


def math_xml(doc):
    return [node_signature(node) for node in doc.element.body.xpath(".//m:oMath")]


def transformed_source_text(i, text):
    text = re.sub(r"\bkhử trùng\b", "loại bỏ trùng lặp", text, flags=re.IGNORECASE)
    if i in {33, 50, 56, 66, 79, 87}:
        text = re.sub(r"^Hình\s+\d+\s*:\s*", "", text).strip()
        return text if text.endswith(".") else text + "."
    if i in {72, 74, 90}:
        text = re.sub(r"^Bảng\s+\d+\s*:\s*", "", text).strip()
        return text if text.endswith(".") else text + "."
    if i in {8, 19, 30, 68, 93, 96}:
        return re.sub(r"^\d+\s+", "", text).strip()
    if i in {98, 100}:
        return text.strip()
    if re.match(r"^\d+\.\d+\s+", text):
        return re.sub(r"^\d+\.\d+\s+", "", text).strip()
    if i >= 101:
        return re.sub(r"^\[\d+\]\s*", "", text)
    return text


def main():
    if sha256(TEMPLATE.read_bytes()) != EXPECTED_TEMPLATE_SHA:
        raise AssertionError("Retained template changed")

    template = Document(str(TEMPLATE))
    source = Document(str(SOURCE))
    final = Document(str(FINAL))

    assert len(final.sections) == 1
    ts, fs = template.sections[0], final.sections[0]
    for attr in (
        "page_width", "page_height", "top_margin", "bottom_margin",
        "left_margin", "right_margin", "header_distance", "footer_distance",
    ):
        assert int(getattr(ts, attr)) == int(getattr(fs, attr)), attr
    cols = fs._sectPr.find(qn("w:cols"))
    assert cols is None or cols.get(qn("w:num")) in (None, "1")

    style_names = {s.name for s in final.styles}
    missing_styles = [s for s in REQUIRED_STYLES if s not in style_names]
    assert not missing_styles, missing_styles

    paragraph_style_counts = Counter(p.style.name for p in final.paragraphs)
    assert paragraph_style_counts["TenBaiBao"] == 1
    assert paragraph_style_counts["TacGia"] == 1
    assert paragraph_style_counts["TomTat"] == 2
    assert paragraph_style_counts["TuKhoa"] == 1
    assert paragraph_style_counts["MucChinh"] == 8
    assert paragraph_style_counts["TenBang"] == 3
    assert paragraph_style_counts["TenHinh"] == 6
    assert paragraph_style_counts["TaiLieuThamKhao"] == 18

    assert len(final.tables) == 3
    assert len(final.inline_shapes) == 7
    assert len(final.element.body.xpath(".//m:oMath")) == 7
    assert len(pdfium.PdfDocument(str(PDF))) == 10

    abstract_candidates = [p.text for p in final.paragraphs if p.style.name == "TomTat" and p.text != "Tóm tắt"]
    assert len(abstract_candidates) == 1
    assert len(abstract_candidates[0].split()) == 183
    keyword_text = [p.text for p in final.paragraphs if p.style.name == "TuKhoa"]
    assert keyword_text == ["Từ khóa: ByteTrack; Giám sát biên; Nhận dạng hành vi; Theo dõi đa đối tượng."]

    all_text = "\n".join(p.text for p in final.paragraphs)
    forbidden = ["CCS CONCEPTS", "ACM Reference Format", "Creative Commons", "ISBN", "khử trùng"]
    assert not [term for term in forbidden if term in all_text]
    required_tokens = [
        "YOLOv11", "ByteTrack", "InsightFace", "risk_score", "gate_alerts",
        "pending_person", "known_person", "pose_push_contact", "score_threshold",
        "scripts/train_behavior_classifier.py", "/api/behavior-events", "20.76",
        "24.47", "135", "130",
    ]
    assert not [term for term in required_tokens if term not in all_text]
    assert all_text.count("loại bỏ trùng lặp") >= 3

    for src_table, out_table in zip(source.tables[1:], final.tables):
        src_data = [[cell.text for cell in row.cells] for row in src_table.rows]
        out_data = [[cell.text for cell in row.cells] for row in out_table.rows]
        assert src_data == out_data
        for ci, cell in enumerate(out_table.rows[0].cells):
            assert all(p.style.name == "TieuDeBang" for p in cell.paragraphs), ci
        for row in out_table.rows[1:]:
            for cell in row.cells:
                assert all(p.style.name == "NoiDungBang" for p in cell.paragraphs)

    assert related_image_hashes(source) == related_image_hashes(final)
    assert math_xml(source) == math_xml(final)

    skip = {32, 38, 40, 44, 49, 52, 55, 58, 60, 62, 65, 78, 86, 91}
    final_text_counts = Counter(p.text for p in final.paragraphs)
    missing = []
    for i in range(8, len(source.paragraphs)):
        if i in skip:
            continue
        expected = transformed_source_text(i, source.paragraphs[i].text)
        if expected and final_text_counts[expected] == 0:
            missing.append((i, expected[:80]))
    assert not missing, missing

    ref_paras = [p for p in final.paragraphs if p.style.name == "TaiLieuThamKhao"]
    assert len(ref_paras) == 18
    for n, p in enumerate(ref_paras, 1):
        assert not p.text.startswith(f"[{n}]")
        num_id = p._p.pPr.numPr.numId.val if p._p.pPr is not None and p._p.pPr.numPr is not None else None
        assert num_id is None or int(num_id) != 0

    print("PASS")
    print("pages=10 sections=1 columns=1 A4=true")
    print("styles=required-present abstract_words=183 keywords=4 ccs=absent")
    print("images=7 tables=3 equations=7 references=18")
    print("source_content_coverage=complete table_data=exact image_blobs=exact omml=exact")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
