from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.dml.color import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def replace_run_text(run_element, text: str, *, line_break_before: bool = False) -> None:
    for child in list(run_element):
        if child.tag != qn("w:rPr"):
            run_element.remove(child)
    if line_break_before:
        run_element.append(OxmlElement("w:br"))
    text_element = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run_element.append(text_element)


def main() -> None:
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)
    formatted = []

    for paragraph in document.paragraphs:
        if not text_of(paragraph._p).startswith("STACK"):
            continue

        project_link = None
        for hyperlink in paragraph._p.iterchildren(qn("w:hyperlink")):
            rel_id = hyperlink.get(qn("r:id"))
            if not rel_id or rel_id not in document.part.rels:
                continue
            target = document.part.rels[rel_id].target_ref
            if target.startswith("https://github.com/"):
                project_link = (hyperlink, target)
                break
        if project_link is None:
            continue

        hyperlink, target = project_link
        direct_runs = list(paragraph._p.iterchildren(qn("w:r")))
        separator_run = next(
            (run for run in direct_runs if re.search(r"\s\|\s*$", text_of(run))),
            None,
        )
        label_run = next(
            (run for run in direct_runs if text_of(run).strip() == "GITHUB"),
            None,
        )
        if separator_run is None or label_run is None:
            raise RuntimeError(f"Could not find project metadata fields for {target}")

        replace_run_text(separator_run, re.sub(r"\s*\|\s*$", "", text_of(separator_run)))
        replace_run_text(label_run, "GITHUB  ", line_break_before=True)

        hyperlink_text = target.removeprefix("https://")
        text_nodes = list(hyperlink.iter(qn("w:t")))
        if not text_nodes:
            raise RuntimeError(f"Hyperlink has no display text: {target}")
        text_nodes[0].text = hyperlink_text
        for node in text_nodes[1:]:
            node.text = ""

        hyperlink_run_element = next(hyperlink.iter(qn("w:r")))
        hyperlink_run = Run(hyperlink_run_element, paragraph)
        hyperlink_run.bold = False
        hyperlink_run.underline = True
        hyperlink_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        seen_hyperlink = False
        for child in list(paragraph._p):
            if child is hyperlink:
                seen_hyperlink = True
                continue
            if seen_hyperlink and child.tag == qn("w:r") and not text_of(child).strip():
                paragraph._p.remove(child)

        formatted.append(target)

    if len(formatted) != 2:
        raise RuntimeError(f"Expected 2 project links, formatted {len(formatted)}: {formatted}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    check = Document(output)
    verified = []
    for paragraph in check.paragraphs:
        if not text_of(paragraph._p).startswith("STACK"):
            continue
        if len(list(paragraph._p.iter(qn("w:br")))) != 1:
            raise RuntimeError(f"Expected one explicit line break: {text_of(paragraph._p)}")
        for hyperlink in paragraph._p.iter(qn("w:hyperlink")):
            rel_id = hyperlink.get(qn("r:id"))
            if not rel_id or rel_id not in check.part.rels:
                continue
            target = check.part.rels[rel_id].target_ref
            if target in formatted:
                display = text_of(hyperlink)
                if display != target.removeprefix("https://"):
                    raise RuntimeError(f"Display/target mismatch: {display!r} -> {target!r}")
                verified.append((display, target))

    if len(verified) != 2:
        raise RuntimeError(f"Hyperlink verification failed: {verified}")
    for display, target in verified:
        print(f"OK: {display} -> {target}")


if __name__ == "__main__":
    main()
