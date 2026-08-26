from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def wval(element, attr="val"):
    return element.get(qn(f"w:{attr}")) if element is not None else None


def paragraph_num_pr(paragraph):
    p_pr = paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    if num_pr is not None:
        return num_pr
    style = paragraph.style
    while style is not None:
        style_p_pr = style.element.pPr
        num_pr = style_p_pr.numPr if style_p_pr is not None else None
        if num_pr is not None:
            return num_pr
        style = style.base_style
    return None


def numbering_format(document, paragraph):
    num_pr = paragraph_num_pr(paragraph)
    if num_pr is None or num_pr.numId is None:
        return None
    num_id = wval(num_pr.numId)
    ilvl = wval(num_pr.ilvl) or "0"
    root = document.part.numbering_part.element
    num = root.find(f"./w:num[@w:numId='{num_id}']", root.nsmap)
    if num is None:
        return num_id, ilvl, None, None
    abstract_id = wval(num.find("./w:abstractNumId", root.nsmap))
    abstract = root.find(f"./w:abstractNum[@w:abstractNumId='{abstract_id}']", root.nsmap)
    level = abstract.find(f"./w:lvl[@w:ilvl='{ilvl}']", root.nsmap) if abstract is not None else None
    if level is None and abstract is not None:
        level = abstract.find("./w:lvl", root.nsmap)
    fmt = wval(level.find("./w:numFmt", root.nsmap)) if level is not None else None
    marker = wval(level.find("./w:lvlText", root.nsmap)) if level is not None else None
    return num_id, ilvl, fmt, marker


def walk(document):
    for idx, paragraph in enumerate(document.paragraphs):
        yield f"body:{idx}", paragraph
    for t_idx, table in enumerate(document.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, paragraph in enumerate(cell.paragraphs):
                    yield f"table:{t_idx}/{r_idx}/{c_idx}/{p_idx}", paragraph


def bold_prefix(paragraph):
    pieces = []
    started = False
    for run in paragraph.runs:
        if not run.text:
            continue
        if run.bold is True:
            pieces.append(run.text)
            started = True
        elif started:
            break
        elif run.text.strip():
            return ""
    return "".join(pieces).strip()


def main():
    path = Path(sys.argv[1])
    document = Document(path)
    counts = {"bullet": 0, "candidate": 0}
    for loc, paragraph in walk(document):
        text = paragraph.text.strip()
        info = numbering_format(document, paragraph)
        if not text or not info or info[2] != "bullet":
            continue
        counts["bullet"] += 1
        prefix = bold_prefix(paragraph)
        if not prefix:
            continue
        counts["candidate"] += 1
        run_sizes = [run.font.size.pt if run.font.size else None for run in paragraph.runs if run.text]
        print(f"{loc}\t{info}\tbold_prefix={prefix!r}\tsizes={run_sizes}\ttext={text!r}")
    print(f"COUNTS\t{counts}")


if __name__ == "__main__":
    main()
