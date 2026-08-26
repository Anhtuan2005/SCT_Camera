from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def value(element, attr="val"):
    return element.get(qn(f"w:{attr}")) if element is not None else None


def numbering_info(paragraph):
    p_pr = paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    source = "direct"
    if num_pr is None and paragraph.style is not None:
        style_p_pr = paragraph.style.element.pPr
        num_pr = style_p_pr.numPr if style_p_pr is not None else None
        source = "style"
    if num_pr is None:
        return None
    num_id = value(num_pr.numId)
    ilvl = value(num_pr.ilvl) or "0"
    return source, num_id, ilvl


def run_summary(paragraph):
    parts = []
    for run in paragraph.runs:
        text = run.text.replace("\t", "\\t").replace("\n", "\\n")
        if not text:
            continue
        size = run.font.size.pt if run.font.size else None
        parts.append(
            {
                "text": text,
                "size": size,
                "bold": run.bold,
                "italic": run.italic,
                "style": run.style.name if run.style else None,
            }
        )
    return parts


def walk_paragraphs(document):
    for idx, paragraph in enumerate(document.paragraphs):
        yield f"body:{idx}", paragraph
    for t_idx, table in enumerate(document.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, paragraph in enumerate(cell.paragraphs):
                    yield f"table:{t_idx}/{r_idx}/{c_idx}/{p_idx}", paragraph
    for s_idx, section in enumerate(document.sections):
        for kind, part in (("header", section.header), ("footer", section.footer)):
            for p_idx, paragraph in enumerate(part.paragraphs):
                yield f"{kind}:{s_idx}/{p_idx}", paragraph


def main():
    path = Path(sys.argv[1])
    document = Document(path)
    print(f"FILE\t{path}")
    print(f"BODY_PARAGRAPHS\t{len(document.paragraphs)}")
    print(f"TABLES\t{len(document.tables)}")
    for loc, paragraph in walk_paragraphs(document):
        text = paragraph.text.strip()
        num = numbering_info(paragraph)
        manual_bullet = text.startswith(("•", "●", "◦", "▪", "-", "–"))
        if not text or (num is None and not manual_bullet):
            continue
        print("PARA", loc, paragraph.style.name if paragraph.style else "", num, repr(text), sep="\t")
        for item in run_summary(paragraph):
            print("  RUN", repr(item), sep="\t")


if __name__ == "__main__":
    main()
