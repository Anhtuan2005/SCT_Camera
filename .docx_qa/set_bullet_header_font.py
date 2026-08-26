from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def wval(element, attr="val"):
    return element.get(qn(f"w:{attr}")) if element is not None else None


def paragraph_num_pr(paragraph):
    p_pr = paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    if num_pr is not None:
        return num_pr
    style = paragraph.style
    while style is not None:
        style_p_pr = style.element.pPr
        num_pr = style_p_pr.numPr if style_p_pr is not None else None
        if num_pr is not None:
            return num_pr
        style = style.base_style
    return None


def is_bullet(document, paragraph):
    num_pr = paragraph_num_pr(paragraph)
    if num_pr is None or num_pr.numId is None:
        return False
    num_id = wval(num_pr.numId)
    ilvl = wval(num_pr.ilvl) or "0"
    root = document.part.numbering_part.element
    num = root.find(f"./w:num[@w:numId='{num_id}']", root.nsmap)
    if num is None:
        return False
    abstract_id = wval(num.find("./w:abstractNumId", root.nsmap))
    abstract = root.find(f"./w:abstractNum[@w:abstractNumId='{abstract_id}']", root.nsmap)
    if abstract is None:
        return False
    level = abstract.find(f"./w:lvl[@w:ilvl='{ilvl}']", root.nsmap)
    if level is None:
        level = abstract.find("./w:lvl", root.nsmap)
    num_fmt = level.find("./w:numFmt", root.nsmap) if level is not None else None
    return wval(num_fmt) == "bullet"


def walk(document):
    for idx, paragraph in enumerate(document.paragraphs):
        yield f"body:{idx}", paragraph
    for t_idx, table in enumerate(document.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, paragraph in enumerate(cell.paragraphs):
                    yield f"table:{t_idx}/{r_idx}/{c_idx}/{p_idx}", paragraph


def set_header_runs(paragraph):
    changed_runs = 0
    header_text = []
    started = False
    for run in paragraph.runs:
        if not run.text:
            continue
        if run.bold is True:
            run.font.size = Pt(13)
            header_text.append(run.text)
            changed_runs += 1
            started = True
        elif started:
            break
        elif run.text.strip():
            return 0, ""
    return changed_runs, "".join(header_text).strip()


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: set_bullet_header_font.py INPUT.docx OUTPUT.docx")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    document = Document(source)
    changed = []
    for loc, paragraph in walk(document):
        if not paragraph.text.strip() or not is_bullet(document, paragraph):
            continue
        run_count, header = set_header_runs(paragraph)
        if run_count:
            changed.append((loc, header, run_count))
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"OUTPUT\t{output}")
    print(f"BULLET_HEADERS\t{len(changed)}")
    print(f"RUNS_SET_TO_13PT\t{sum(item[2] for item in changed)}")
    for loc, header, run_count in changed:
        print(f"{loc}\t{run_count}\t{header}")


if __name__ == "__main__":
    main()
