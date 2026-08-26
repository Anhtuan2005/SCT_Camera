import re
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt


TEMPLATE = Path(r"E:\SCT_Camera\.codex_work\faculty_template\faculty_template.docx")
SOURCE = Path(r"E:\Paper_NCKH.docx")
OUTPUT = Path(r"E:\SCT_Camera\.codex_work\faculty_template\paper_faculty_v3.docx")

ABSTRACT = (
    "Giám sát video quy mô nhỏ cần vượt ra ngoài ghi hình kích hoạt theo chuyển động để theo "
    "dõi đa camera, nhận biết danh tính và phân tích hành vi trên phần cứng phổ thông. Bài báo "
    "trình bày SCT Camera, kiến trúc sáu giai đoạn gần "
    "thời gian thực gồm thu nhận, thị giác AI, phân giải danh tính, phân tích hành vi, học/cổng "
    "rủi ro và phân phối cảnh báo. Hệ thống kết hợp YOLOv11; ByteTrack có tái liên kết vết "
    "mất và bù chuyển động camera; InsightFace; bảy luật hành vi dựa trên hình học; và mô hình "
    "hồi quy logistic 22 đặc trưng học từ sự kiện do người vận hành gán nhãn. Cảnh báo được "
    "phân phối bất đồng bộ đa kênh. Trên RTX 3050 Laptop 4 GB, hệ thống đạt 20.76 "
    "FPS với một camera và tổng 24.47 FPS với hai camera; 130 trong 135 ca kiểm thử đạt. Kết "
    "quả cho thấy kiến trúc đáp ứng xử lý gần thời gian thực, duy trì tính giải thích và giảm "
    "cảnh báo độ tin cậy thấp."
)
KEYWORDS = "ByteTrack; Giám sát biên; Nhận dạng hành vi; Theo dõi đa đối tượng."


def clear_content(p_el):
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)


def set_run_font_inheritance(run):
    rpr = run._r.rPr
    if rpr is not None:
        for tag in ("w:rFonts", "w:sz", "w:szCs", "w:color"):
            node = rpr.find(qn(tag))
            if node is not None:
                rpr.remove(node)


def append_pattern(doc, pattern_el):
    el = deepcopy(pattern_el)
    clear_content(el)
    body = doc.element.body
    body.insert(len(body) - 1, el)
    return Paragraph(el, doc._body)


def add_text_pattern(doc, pattern_el, text, *, bold=None, italic=None):
    p = append_pattern(doc, pattern_el)
    run = p.add_run(text)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    set_run_font_inheritance(run)
    return p


def suppress_numbering(p):
    ppr = p._p.get_or_add_pPr()
    num_pr = ppr.get_or_add_numPr()
    num_id = num_pr.get_or_add_numId()
    num_id.val = 0


def set_paragraph_size(p, points):
    for run in p.runs:
        run.font.size = Pt(points)


def add_author_block(doc, patterns):
    p = append_pattern(doc, patterns["authors"])
    for name, marker, tail in [
        ("Hanh Nguyen", "a*", ", "),
        ("Nguyen Nguyen", "a", ", "),
        ("Tuan Nguyen", "a", ""),
    ]:
        r = p.add_run(name)
        set_run_font_inheritance(r)
        r = p.add_run(marker)
        r.font.superscript = True
        set_run_font_inheritance(r)
        if tail:
            r = p.add_run(tail)
            set_run_font_inheritance(r)

    p = append_pattern(doc, patterns["affiliation"])
    r = p.add_run("a")
    r.font.superscript = True
    set_run_font_inheritance(r)
    r = p.add_run(" Khoa Công nghệ Thông tin, Trường Đại học Ngoại ngữ – Tin học Thành phố Hồ Chí Minh")
    set_run_font_inheritance(r)

    p = add_text_pattern(
        doc,
        patterns["corresponding"],
        "* Tác giả liên hệ: Email: hanhntm@huflit.edu.vn | Điện thoại: [Bổ sung sau]",
    )
    p.paragraph_format.keep_with_next = True


def add_front_matter(doc, patterns):
    title = add_text_pattern(
        doc,
        patterns["title"],
        "PHÁT TRIỂN HỆ THỐNG CẢNH BÁO THỜI GIAN THỰC DỰA TRÊN AI CHO CAMERA GIÁM SÁT",
    )
    set_paragraph_size(title, 17.5)
    add_author_block(doc, patterns)
    abstract_label = add_text_pattern(doc, patterns["abstract_label"], "Tóm tắt", bold=True, italic=False)
    abstract_body = add_text_pattern(doc, patterns["abstract_body"], ABSTRACT)
    set_paragraph_size(abstract_label, 11.5)
    set_paragraph_size(abstract_body, 11.5)
    p = append_pattern(doc, patterns["keywords"])
    r = p.add_run("Từ khóa: ")
    r.bold = True
    set_run_font_inheritance(r)
    r = p.add_run(KEYWORDS)
    set_run_font_inheritance(r)
    set_paragraph_size(p, 11.5)


def clean_body_text(text):
    return re.sub(r"\bkhử trùng\b", "loại bỏ trùng lặp", text, flags=re.IGNORECASE)


def figure_description(text):
    text = re.sub(r"^Hình\s+\d+\s*:\s*", "", text).strip()
    return text if text.endswith(".") else text + "."


def table_description(text):
    text = re.sub(r"^Bảng\s+\d+\s*:\s*", "", text).strip()
    return text if text.endswith(".") else text + "."


def image_parts(source, src_paragraph):
    result = []
    for blip in src_paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        part = source.part.related_parts[rel_id]
        result.append(part.blob)
    return result


def add_figure(doc, patterns, source, src_index):
    p = append_pattern(doc, patterns["image"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    blobs = image_parts(source, source.paragraphs[src_index])
    if src_index == 86:
        for i, blob in enumerate(blobs):
            run = p.add_run()
            run.add_picture(BytesIO(blob), height=Inches(3.00))
            if i != len(blobs) - 1:
                p.add_run("  ")
        return
    widths = {32: 5.20, 49: 4.80, 55: 4.70, 65: 4.80, 78: 4.15}
    run = p.add_run()
    run.add_picture(BytesIO(blobs[0]), width=Inches(widths[src_index]))


def add_equation(doc, pattern_el, src_p, number):
    el = deepcopy(pattern_el)
    target_math = el.xpath(".//m:oMath")
    source_math = src_p._p.xpath(".//m:oMath")
    if len(target_math) != 1 or len(source_math) != 1:
        raise ValueError(f"Unexpected OMML shape for equation {number}")
    target_math[0].getparent().replace(target_math[0], deepcopy(source_math[0]))
    number_nodes = [
        t for t in el.xpath(".//w:t")
        if t.text and re.fullmatch(r"\(\d+\)", t.text.strip())
    ]
    if len(number_nodes) != 1:
        raise ValueError(f"Equation pattern number not found for equation {number}")
    number_nodes[0].text = f"({number})"
    body = doc.element.body
    body.insert(len(body) - 1, el)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_inches):
    widths_dxa = [round(w * 1440) for w in widths_inches]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for ci, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[ci]))


def set_cell_margins(cell, top=30, start=75, bottom=30, end=75):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_data_table(doc, source_table, index):
    widths = {
        1: [1.05, 1.00, 1.10, 1.45, 1.45],
        2: [2.20, 1.25, 1.25, 1.25],
        3: [1.75, 1.35, 3.05],
    }[index]
    table = doc.add_table(rows=len(source_table.rows), cols=len(source_table.columns))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_header(table.rows[0])
    for ri, src_row in enumerate(source_table.rows):
        for ci, src_cell in enumerate(src_row.cells):
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.style = "TieuDeBang" if ri == 0 else "NoiDungBang"
            p.text = src_cell.text
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif index == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif index == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            else:
                if ci == 0 or (ri == 3 and ci == 2):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_run_font_inheritance(run)
            set_paragraph_size(p, 10.5)
    return table


def build():
    template = Document(str(TEMPLATE))
    source = Document(str(SOURCE))
    tp = template.paragraphs
    patterns = {
        "title": deepcopy(tp[0]._p),
        "authors": deepcopy(tp[1]._p),
        "affiliation": deepcopy(tp[2]._p),
        "corresponding": deepcopy(tp[4]._p),
        "abstract_label": deepcopy(tp[5]._p),
        "abstract_body": deepcopy(tp[6]._p),
        "keywords": deepcopy(tp[7]._p),
        "main": deepcopy(tp[8]._p),
        "body": deepcopy(tp[9]._p),
        "image": deepcopy(tp[24]._p),
        "sub1": deepcopy(tp[27]._p),
        "sub2": deepcopy(tp[31]._p),
        "equation": deepcopy(tp[37]._p),
        "table_caption": deepcopy(tp[49]._p),
        "figure_caption": deepcopy(tp[51]._p),
        "reference": deepcopy(tp[60]._p),
    }

    body = template.element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    add_front_matter(template, patterns)

    image_paragraphs = {32, 49, 55, 65, 78, 86}
    figure_captions = {33, 50, 56, 66, 79, 87}
    equations = {38: 1, 40: 2, 44: 3, 52: 4, 58: 5, 60: 6, 62: 7}
    table_captions = {72: 1, 74: 2, 90: 3}

    for i in range(8, len(source.paragraphs)):
        src_p = source.paragraphs[i]
        text = clean_body_text(src_p.text)

        if i in image_paragraphs:
            add_figure(template, patterns, source, i)
            continue
        if i in figure_captions:
            p = add_text_pattern(template, patterns["figure_caption"], figure_description(text))
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            set_paragraph_size(p, 11.5)
            continue
        if i in equations:
            add_equation(template, patterns["equation"], src_p, equations[i])
            continue
        if i in table_captions:
            table_index = table_captions[i]
            p = add_text_pattern(template, patterns["table_caption"], table_description(text))
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = True
            set_paragraph_size(p, 11.5)
            add_data_table(template, source.tables[table_index], table_index)
            continue
        if i == 91 and not text.strip():
            continue

        if src_p.style.name == "Heading 1":
            heading_text = re.sub(r"^\d+\s+", "", text).strip()
            p = add_text_pattern(template, patterns["main"], heading_text)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            set_paragraph_size(p, 12)
            if i in (98, 100):
                suppress_numbering(p)
            continue
        if src_p.style.name == "Heading 2":
            heading_text = re.sub(r"^\d+\.\d+\s+", "", text).strip()
            p = add_text_pattern(template, patterns["sub1"], heading_text)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            set_paragraph_size(p, 12)
            continue

        if i >= 101:
            ref_text = re.sub(r"^\[\d+\]\s*", "", text)
            p = add_text_pattern(template, patterns["reference"], ref_text)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            set_paragraph_size(p, 10.5)
            continue

        p = add_text_pattern(template, patterns["body"], text)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_paragraph_size(p, 11.5)
        if i in (39, 41):
            p.paragraph_format.first_line_indent = 0

    core = template.core_properties
    core.title = "Phát triển hệ thống cảnh báo thời gian thực dựa trên AI cho camera giám sát"
    core.subject = "Bài báo nghiên cứu khoa học – định dạng Khoa Công nghệ Thông tin"
    core.keywords = KEYWORDS.rstrip(".")
    core.comments = "Định dạng trực tiếp từ template của Khoa CNTT."

    word_count = len(ABSTRACT.split())
    if not 150 <= word_count <= 200:
        raise ValueError(f"Abstract word count is {word_count}, expected 150–200")
    if len(KEYWORDS.rstrip(".").split("; ")) != 4:
        raise ValueError("Expected exactly four keywords")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    template.save(str(OUTPUT))
    print(f"output={OUTPUT}")
    print(f"abstract_words={word_count}")
    print("images=7 tables=3 equations=7 references=18")


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
