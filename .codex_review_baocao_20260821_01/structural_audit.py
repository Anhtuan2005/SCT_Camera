from __future__ import annotations

import collections
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


DOCX = Path(r"E:\BAOCAO.docx")
OUT = Path(r"E:\SCT_Camera\.codex_review_baocao_20260821_01\structural_audit.json")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def run_value(run, attr):
    value = getattr(run.font, attr)
    if value is not None:
        return value
    style = run.style
    if style is not None:
        value = getattr(style.font, attr)
        if value is not None:
            return value
    return getattr(run._parent.style.font, attr)


def fmt_size(run):
    value = run_value(run, "size")
    return round(value.pt, 2) if value else None


def fmt_name(run):
    return run_value(run, "name")


doc = Document(DOCX)
paragraphs = []
captions = []
font_chars = collections.Counter()
size_chars = collections.Counter()
style_chars = collections.Counter()
duplicate_text = collections.defaultdict(list)

for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        duplicate_text[text].append(index)
    style = paragraph.style.name if paragraph.style else None
    style_chars[style] += len(text)
    runs = []
    for run_index, run in enumerate(paragraph.runs):
        if not run.text:
            continue
        font = fmt_name(run)
        size = fmt_size(run)
        bold = run_value(run, "bold")
        italic = run_value(run, "italic")
        font_chars[font] += len(run.text)
        size_chars[size] += len(run.text)
        runs.append(
            {
                "index": run_index,
                "text": run.text,
                "font": font,
                "size_pt": size,
                "bold": bold,
                "italic": italic,
            }
        )
    entry = {
        "index": index,
        "style": style,
        "text": text,
        "runs": runs,
        "field_instructions": [
            node.text or "" for node in paragraph._p.iter(f"{W}instrText")
        ],
    }
    paragraphs.append(entry)
    if re.match(r"^(Hình|Bảng)\s+\d+-\d+\s*:", text, re.IGNORECASE):
        captions.append(entry)

with zipfile.ZipFile(DOCX) as archive:
    names = set(archive.namelist())
    xml_parts = {
        name: etree.fromstring(archive.read(name))
        for name in names
        if name.startswith("word/") and name.endswith(".xml")
    }
    tracked_insertions = sum(len(root.findall(f".//{W}ins")) for root in xml_parts.values())
    tracked_deletions = sum(len(root.findall(f".//{W}del")) for root in xml_parts.values())
    comments = 0
    if "word/comments.xml" in xml_parts:
        comments = len(xml_parts["word/comments.xml"].findall(f".//{W}comment"))
    instr_text = " ".join(
        node.text or ""
        for root in xml_parts.values()
        for node in root.findall(f".//{W}instrText")
    )
    toc_fields = len(re.findall(r"\bTOC\b", instr_text, re.IGNORECASE))
    seq_fields = re.findall(r"\bSEQ\s+(Hình|Bảng)\b", instr_text, re.IGNORECASE)
    page_breaks = len(xml_parts["word/document.xml"].findall(f".//{W}br[@{W}type='page']"))
    last_rendered_breaks = len(
        xml_parts["word/document.xml"].findall(f".//{W}lastRenderedPageBreak")
    )

table_details = []
for table_index, table in enumerate(doc.tables):
    font_counter = collections.Counter()
    size_counter = collections.Counter()
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    font_counter[fmt_name(run)] += len(run.text)
                    size_counter[fmt_size(run)] += len(run.text)
    first_row_header = bool(
        table.rows
        and table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    )
    table_details.append(
        {
            "index": table_index,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "style": table.style.name if table.style else None,
            "first_row_header": first_row_header,
            "top_fonts_by_chars": font_counter.most_common(10),
            "top_sizes_by_chars": size_counter.most_common(10),
            "first_row_text": [cell.text.strip() for cell in table.rows[0].cells]
            if table.rows
            else [],
        }
    )

result = {
    "core_properties": {
        "title": doc.core_properties.title,
        "subject": doc.core_properties.subject,
        "author": doc.core_properties.author,
        "last_modified_by": doc.core_properties.last_modified_by,
        "created": str(doc.core_properties.created),
        "modified": str(doc.core_properties.modified),
    },
    "counts": {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "sections": len(doc.sections),
        "captions_detected": len(captions),
        "tracked_insertions": tracked_insertions,
        "tracked_deletions": tracked_deletions,
        "comments": comments,
        "toc_fields": toc_fields,
        "seq_fields": collections.Counter(x.lower() for x in seq_fields),
        "explicit_page_breaks": page_breaks,
        "last_rendered_page_breaks": last_rendered_breaks,
    },
    "table_shapes": table_details,
    "top_fonts_by_chars": font_chars.most_common(20),
    "top_sizes_by_chars": size_chars.most_common(20),
    "top_styles_by_chars": style_chars.most_common(20),
    "captions": captions,
    "duplicate_paragraphs": [
        {"text": text, "indices": indices}
        for text, indices in duplicate_text.items()
        if len(indices) > 1 and len(text) >= 20
    ],
    "paragraphs": paragraphs,
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=list), encoding="utf-8")
print(json.dumps({k: result[k] for k in ("core_properties", "counts", "table_shapes", "top_fonts_by_chars", "top_sizes_by_chars", "duplicate_paragraphs")}, ensure_ascii=False, indent=2, default=list))
