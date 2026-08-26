from __future__ import annotations

import copy
import hashlib
import sys
import zipfile
from pathlib import Path

import pypdf
from docx import Document
from lxml import etree

from compact_layout_no_font_resize import NS, W, patch_document


def parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path, "r") as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def c14n(xml: bytes) -> bytes:
    return etree.tostring(etree.fromstring(xml), method="c14n")


def remove_direct_sizes(xml: bytes) -> tuple[bytes, int]:
    root = etree.fromstring(xml, etree.XMLParser(remove_blank_text=False))
    nodes = root.xpath(".//w:rPr/w:sz | .//w:rPr/w:szCs", namespaces=NS)
    for node in nodes:
        node.getparent().remove(node)
    return (
        etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        len(nodes),
    )


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: verify_style_inheritance.py SOURCE.docx FINAL.docx RENDER.pdf")
    source = Path(sys.argv[1]).resolve()
    final = Path(sys.argv[2]).resolve()
    render_pdf = Path(sys.argv[3]).resolve()

    source_parts = parts(source)
    final_parts = parts(final)
    assert set(source_parts) == set(final_parts)
    changed_parts = [
        name for name in source_parts if source_parts[name] != final_parts[name]
    ]
    assert changed_parts == ["word/document.xml"], changed_parts
    assert source_parts["word/styles.xml"] == final_parts["word/styles.xml"]

    no_sizes_xml, removed = remove_direct_sizes(source_parts["word/document.xml"])
    assert removed == 155, removed
    expected_xml, report = patch_document(no_sizes_xml, 1.75, 1.5)
    assert report["figure_captions_unlinked"] == 6
    assert report["references_compacted"] == 18
    assert c14n(expected_xml) == c14n(final_parts["word/document.xml"])

    direct_size_count = 0
    for name, data in final_parts.items():
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        if name in {"word/styles.xml", "word/numbering.xml"}:
            continue
        root = etree.fromstring(data)
        direct_size_count += len(
            root.xpath(".//w:rPr/w:sz | .//w:rPr/w:szCs", namespaces=NS)
        )
    assert direct_size_count == 0

    source_doc = Document(str(source))
    final_doc = Document(str(final))
    assert [p.text for p in source_doc.paragraphs] == [p.text for p in final_doc.paragraphs]
    assert len(source_doc.tables) == len(final_doc.tables) == 3
    for source_table, final_table in zip(source_doc.tables, final_doc.tables, strict=True):
        assert [
            [cell.text for cell in row.cells] for row in source_table.rows
        ] == [
            [cell.text for cell in row.cells] for row in final_table.rows
        ]
    assert len(final_doc.inline_shapes) == 7
    assert len(final_doc.element.body.xpath(".//m:oMath")) == 7

    style_sizes = {
        name: (final_doc.styles[name].font.size.pt if final_doc.styles[name].font.size else None)
        for name in (
            "TenBaiBao",
            "MucChinh",
            "MucPhu1",
            "MucPhu2",
            "TenBang",
            "TenHinh",
            "TieuDeBang",
            "TaiLieuThamKhao",
        )
    }
    assert style_sizes == {
        "TenBaiBao": 24.0,
        "MucChinh": 12.0,
        "MucPhu1": 13.0,
        "MucPhu2": 13.0,
        "TenBang": 13.0,
        "TenHinh": 13.0,
        "TieuDeBang": 13.0,
        "TaiLieuThamKhao": 13.0,
    }

    reader = pypdf.PdfReader(str(render_pdf))
    assert len(reader.pages) == 10
    page_text = [page.extract_text() or "" for page in reader.pages]
    assert "0977314645" in page_text[0]
    assert "Bảng 2." in page_text[6]
    assert all(
        token in page_text[6]
        for token in ("Phát hiện/Theo dõi", "Tư thế", "Danh tính", "Hành vi", "2.45 ms")
    )
    assert "Bảng 3." in page_text[7]
    assert all(
        token in page_text[7]
        for token in ("Thời gian ở gần", "Lượt đi qua lại", "Tín hiệu hành vi", "Đối tượng phát hiện")
    )
    assert "[1]" in page_text[8]
    assert "[18]" in page_text[9]

    print("PASS")
    print("direct_font_size_overrides=0")
    print("styles_xml=unchanged style_sizes=template_original")
    print("content_tables_images_equations=unchanged")
    print("pages=10 table_2_single_page=true table_3_single_page=true")
    print(f"sha256={hashlib.sha256(final.read_bytes()).hexdigest()}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
