from collections import Counter
from pathlib import Path
import re
import sys
import zipfile

import pdfplumber
from docx import Document
from lxml import etree


docx_path = Path(sys.argv[1])
pdf_path = Path(sys.argv[2])
doc = Document(docx_path)


def effective_size(run, paragraph):
    if run.font.size:
        return round(run.font.size.pt, 2)
    if paragraph.style and paragraph.style.font.size:
        return round(paragraph.style.font.size.pt, 2)
    normal = doc.styles["Normal"]
    return round(normal.font.size.pt, 2) if normal.font.size else None


size_chars = Counter()
direct_size_chars = Counter()
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        chars = len(run.text)
        if not chars:
            continue
        size_chars[effective_size(run, paragraph)] += chars
        if run.font.size:
            direct_size_chars[round(run.font.size.pt, 2)] += chars

print("EFFECTIVE_FONT_SIZE_CHARS", size_chars.most_common(15))
print("DIRECT_FONT_SIZE_CHARS", direct_size_chars.most_common(15))

captions = []
manual_caption = []
for index, paragraph in enumerate(doc.paragraphs, 1):
    text = paragraph.text.strip()
    if re.match(r"^(Hình|Bảng)\s+\d", text, re.IGNORECASE):
        xml = paragraph._p.xml
        has_seq = "SEQ " in xml or "SEQ%20" in xml
        captions.append((index, paragraph.style.name if paragraph.style else "", has_seq, text))
        if not has_seq:
            manual_caption.append((index, text))

print("CAPTIONS", len(captions), "MANUAL_CAPTIONS", len(manual_caption))
for row in manual_caption:
    print("MANUAL_CAPTION", row)

with zipfile.ZipFile(docx_path) as archive:
    document_xml = archive.read("word/document.xml")
    root = etree.fromstring(document_xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    print("TRACKED_INS", len(root.xpath(".//w:ins", namespaces=ns)))
    print("TRACKED_DEL", len(root.xpath(".//w:del", namespaces=ns)))
    print("COMMENT_STARTS", len(root.xpath(".//w:commentRangeStart", namespaces=ns)))

page_texts = []
with pdfplumber.open(pdf_path) as pdf:
    for page_no, page in enumerate(pdf.pages, 1):
        text = page.extract_text(x_tolerance=2, y_tolerance=2) or ""
        page_texts.append(text)
        if len(re.sub(r"\s+", "", text)) < 450:
            print("LOW_TEXT_PAGE", page_no, len(re.sub(r"\s+", "", text)), repr(text[:180]))

patterns = {
    "generic_figure_ref": re.compile(r"\bHình\s+\d+(?!\s*[-–])", re.IGNORECASE),
    "generic_table_ref": re.compile(r"\bBảng\s+\d+(?!\s*[-–])", re.IGNORECASE),
}

for page_no, text in enumerate(page_texts, 1):
    for match in patterns["generic_figure_ref"].finditer(text):
        print("GENERIC_FIGURE_REF", page_no, match.group(0), repr(text[max(0, match.start()-80):match.end()+100]))
    for match in patterns["generic_table_ref"].finditer(text):
        print("GENERIC_TABLE_REF", page_no, match.group(0), repr(text[max(0, match.start()-80):match.end()+100]))
    for needle in ["cho nhauĐây", "Hình 20", "Hình 36", "Ảnh hưởng :", "font-size", "YOLO11l", "YOLO11s", "YOLOv11", "YOLO11"]:
        if needle in text:
            print("NEEDLE", page_no, needle, repr(text[max(0, text.index(needle)-90):text.index(needle)+180]))

print("PAGES", len(page_texts))

print("TEXT_PATTERN_AUDIT_START")
text_patterns = {
    "joined_lower_upper": re.compile(r"(?<=[a-zà-ỹ])(?=[A-ZĐ])"),
    "duplicate_word": re.compile(r"\b([A-Za-zÀ-ỹ]+)\s+\1\b", re.IGNORECASE),
    "space_before_punctuation": re.compile(r"\s+[,.!?;:]"),
}
text_hits = {}
for pattern_name, pattern in text_patterns.items():
    rows = []
    for paragraph_index, paragraph in enumerate(doc.paragraphs, 1):
        text_value = paragraph.text.strip()
        if text_value and pattern.search(text_value):
            rows.append((paragraph_index, text_value))
    text_hits[pattern_name] = rows
    print(pattern_name, len(rows))
    for paragraph_index, text_value in rows[:30]:
        print(paragraph_index, repr(text_value[:260]))
for term in ["YOLOv11", "YOLO11", "YOLO11s", "YOLOv11-Pose", "YOLO11n-pose", "testcase", "test case"]:
    count = sum(paragraph.text.count(term) for paragraph in doc.paragraphs)
    print("TERM", repr(term), count)
print("TEXT_PATTERN_AUDIT_END")
