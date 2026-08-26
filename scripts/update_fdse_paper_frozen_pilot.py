from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


EXPERIMENT_VERSION = "priority-pilot-v2-single-participant-2026-08-18"

ABSTRACT = (
    "Giám sát video quy mô nhỏ cần vượt ra ngoài ghi hình kích hoạt theo chuyển động để "
    "theo dõi, nhận biết danh tính và phân tích hành vi trên phần cứng phổ thông. Bài báo "
    "trình bày SCT Camera, kiến trúc sáu giai đoạn kết hợp YOLOv11, ByteTrack có tái liên "
    "kết vết mất và bù chuyển động camera, InsightFace, các luật hành vi hình học và bộ "
    "trích đặc trưng rủi ro. Audit theo code xác định chính xác véc-tơ 22 chiều; tuy nhiên "
    "chưa có nhãn Risk hợp lệ và mô hình đã huấn luyện nên bài báo không báo cáo chỉ số "
    "Risk. Pilot hành vi một người tham gia gồm 20 clip, chia 10 clip phát triển và 10 clip "
    "test đóng băng, với tổng thời lượng xử lý 1.344,240 giây. Trên tập test, lảng vảng và "
    "hành vi đáng ngờ được phát hiện đúng mà không báo giả; xâm nhập và cắt đường phát hiện "
    "được sự kiện mục tiêu nhưng còn sự kiện trùng/báo giả; trộm cắp bỏ sót clip dương và "
    "cảnh báo sai trên clip âm. Kết quả này chỉ là pilot một người, không chứng minh khả năng "
    "khái quát theo người tham gia."
)

CONTRIBUTION_3 = (
    "(3) Bộ máy hành vi được đối chiếu trực tiếp với code đóng băng. Trong phiên bản này, "
    "xâm nhập là cắt đường theo hướng IN; bộ đếm nội bộ theo dõi cả IN/OUT, còn runner thực "
    "nghiệm chỉ instrument sự kiện line_crossing mà không thay đổi cảnh báo production."
)

CONTRIBUTION_4 = (
    "(4) Audit véc-tơ rủi ro 22 chiều theo đúng thứ tự và phép tính trong code. Pipeline hồi "
    "quy logistic được giữ như một cơ chế dự kiến; do chưa có nhãn Risk và model artifact "
    "hợp lệ, gate_alerts được tắt và không có tuyên bố định lượng về Risk."
)

CONTRIBUTION_5 = (
    "(5) Đánh giá pilot tái lập trên 20 clip của một người tham gia, tách 10 clip phát triển "
    "và 10 clip test đóng băng theo repetition/session, công bố cả phát hiện đúng, bỏ sót, "
    "báo giả và sự kiện trùng thay vì chỉ báo cáo ca thành công."
)

RELATED_RISK = (
    "SCT Camera kết hợp luật theo camera ở Mục 3.5 với bộ trích véc-tơ 22 chiều dự kiến dùng "
    "cho hồi quy logistic [15]. Cách tiếp cận ưu tiên tính minh bạch và chi phí thấp; trái "
    "lại, Sultani và cộng sự [16] học xếp hạng các đoạn CCTV bình thường/bất thường. Trong "
    "nghiên cứu này, pipeline Risk chỉ được audit ở mức mã nguồn vì chưa có nhãn/model hợp "
    "lệ, nên không được xem là mô hình đã huấn luyện hay đã xác thực."
)

ARCHITECTURE = (
    "Hình 1 tóm tắt sáu giai đoạn trên luồng riêng của từng camera và hàng đợi sự kiện bất "
    "đồng bộ: Thu nhận đọc khung BGR; Thị giác AI phát hiện/theo dõi; Danh tính phân giải "
    "nhãn; Phân tích áp dụng luật; Học/Risk trích véc-tơ và chỉ chấm điểm khi có model hợp "
    "lệ; Web/Cảnh báo phát luồng và phân phối thông báo. Trong pilot đóng băng, gate_alerts="
    "false nên kết quả hành vi phản ánh trực tiếp các luật runtime."
)

BEHAVIOR_INTRO = (
    "Semantics đóng băng được tóm tắt ở Bảng 1. Điểm quan trọng của audit phiên bản là "
    "_analyze_zones vẫn tồn tại nhưng analyze() không gọi hàm này; vì vậy không mô tả "
    "intrusion production như xâm nhập polygon. Runtime chỉ emit intrusion khi cắt đường IN. "
    "Bộ đếm line nội bộ nhận cả IN/OUT nhưng không emit line_crossing riêng; runner thực "
    "nghiệm chỉ ghi lại trạng thái bộ đếm để đánh giá line_crossing độc lập, không đổi kiến "
    "trúc hay threshold production."
)

RISK_INTRO = (
    "Audit theo code cho thấy thứ tự véc-tơ x đúng như Bảng 2. Các đặc trưng 1–14 và 18–22 "
    "được tạo trong analytics.behavior_learning.extract_behavior_features; đặc trưng chuyển "
    "động 10–12 dùng _motion_metrics; ba cue 15–17 đến từ "
    "analytics.theft_behavior.SuspiciousTheftDetector. Khoảng cách được chuẩn hóa theo đường "
    "chéo khung, diện tích theo diện tích khung; khi huấn luyện, toàn bộ 22 chiều được z-score "
    "bằng mean/std của tập train."
)

RISK_LIMITATION = (
    "Code hỗ trợ ghép JSONL với nhãn CSV và huấn luyện hồi quy logistic, nhưng artifact đóng "
    "băng không có nhãn Risk hợp lệ và không có model .npz. min_risk_score=0,65 là ngưỡng "
    "cấu hình tĩnh, không phải ngưỡng tối ưu trên validation; gate_alerts=false. Vì vậy bài "
    "báo không báo cáo Risk Accuracy, Precision, Recall, F1, ROC-AUC, confusion matrix hay "
    "ROC curve. Các số liệu ở Mục 4.5 thuộc luật hành vi, không phải mô hình Risk."
)

ALERT_DISTRIBUTION = (
    "Luồng camera đẩy cảnh báo thread-safe vào asyncio.Queue giới hạn 1.000 phần tử. "
    "AlertManager bất đồng bộ áp dụng cooldown/dedup theo camera, loại, vùng/đường/vết rồi "
    "gửi Telegram, Discord và/hoặc còi cục bộ (Hình 3). FastAPI cung cấp MJPEG, REST CRUD "
    "camera/vùng/đường/cài đặt và canvas cấu hình trên luồng video."
)

ENVIRONMENT = (
    "Phép đo thông lượng đã có chạy trên Windows 10, CPU 16 luồng, RAM 15,6 GB, NVIDIA "
    "GeForce RTX 3050 Laptop 4 GB và Python 3.10.11; yolo11n.pt/yolo11n-pose.pt dùng đầu "
    "vào 640 px trên CUDA. Pilot hành vi dùng camera cố định, video khoảng 25 FPS, hai điều "
    "kiện daylight/low-light và hai góc frontal-oblique/side-oblique; ROI/line và threshold "
    "được giữ đúng theo phiên bản đóng băng."
)

PILOT_SETUP = (
    f"Tập pilot phiên bản {EXPERIMENT_VERSION} gồm 20 clip của duy nhất P01: mỗi hành vi "
    "intrusion, loitering, suspicious_behavior, theft và line_crossing có hai repetition, "
    "mỗi repetition gồm một clip dương và một clip âm. Repetition r1 tạo tập development "
    "10 clip (666,296 giây); r2 tạo tập test đóng băng 10 clip (677,944 giây); tổng thời "
    "lượng xử lý là 1.344,240 giây. Hai tập được tách theo clip/session-group nhưng không "
    "participant-disjoint; hash SHA-256 video và cấu hình được lưu trước đánh giá."
)

RESULT_SETUP = (
    "Ground truth được gán từ bằng chứng hình ảnh trước khi diễn giải metric. Mỗi hành vi trên "
    "tập test có một sự kiện dương và một clip âm. Một TP phải khớp cửa sổ ground truth; FP "
    "là mọi sự kiện emit không khớp, kể cả sự kiện trùng trên clip dương; FAR là tỷ lệ clip "
    "âm có ít nhất một cảnh báo sai; latency tính từ thời điểm eligible đến TP đầu tiên. Không "
    "retune ROI, line, threshold hay ground truth sau khi xem test."
)

FIGURE_EXPLANATION = (
    "Hình 5 tách số sự kiện TP/FP/FN khỏi trạng thái báo giả của clip âm, tránh trộn đơn vị "
    "event-level với clip-level. Bảng 5 báo cáo kết quả cho từng hành vi; không tính điểm "
    "tổng hợp vì mỗi hành vi chỉ có một sự kiện dương và một clip âm trong test."
)

RESULT_HIGHLIGHTS = (
    "Loitering và suspicious_behavior đều đạt TP=1, FP=0, FN=0 trên test. Intrusion bắt đúng "
    "lần cắt IN và clip OUT âm sạch, nhưng phát thêm năm sự kiện IN trùng trên chính clip "
    "dương. Theft bỏ sót clip dương và phát một cảnh báo sai trên clip âm. Line_crossing bắt "
    "đúng sự kiện OUT nhưng có một sự kiện không khớp sớm trên clip dương và một cảnh báo "
    "OUT sai trên clip âm."
)

DISCUSSION_1 = (
    "Kết quả âm chỉ ra lỗi cần ưu tiên. Với intrusion, TP xuất hiện sau 0,457 giây nhưng năm "
    "event dư làm precision sự kiện còn 0,166667, cho thấy cơ chế loại bỏ sự kiện trùng lặp "
    "theo track ID chưa ổn định. Loitering phát sau ngưỡng dwell đóng băng 20 giây (latency "
    "0,725 giây) và suspicious_behavior phát sau dwell 180 giây (0,186 giây), cả hai đều sạch "
    "trên clip âm. Theft không phát ở clip dương nhưng báo sai lúc 20,034 giây trên clip âm; "
    "evidence ghi near_seconds=10,0, pose_push_contact=1,0 và vehicle_started_moving=0,0, gợi "
    "ý cue tiếp xúc tư thế sai thay vì dịch chuyển xe."
)

REGRESSION_RESULTS = (
    "Đánh giá hồi quy lịch sử trong benchmark cũ thực thi 135 ca trong 22,29 giây: 130 ca "
    "đạt và năm ca chưa đạt, đều thuộc phân giải danh tính. Đây không phải kết quả của suite "
    "hiện tại. Trên code dùng cho lần sửa bài ngày 19/08/2026, toàn bộ suite hiện tại đạt "
    "206/206 ca, 0 thất bại, trong 8,16 giây; còn ba cảnh báo deprecation của Starlette."
)

DISCUSSION_2 = (
    "Line_crossing có TP tại 12,084 giây (latency 4,084 giây), một event không khớp tại "
    "7,226 giây trên clip dương và một false OUT tại 10,278 giây trên clip âm. Đây là pilot "
    "deadline-safe với một người duy nhất: r1/r2 chỉ tách theo repetition/session, không tách "
    "theo người, nên không chứng minh participant-level generalization hay hiệu năng toàn hệ "
    "thống. Ngoài ra, do chưa có Risk labels/model hợp lệ, mọi chỉ số định lượng của Risk vẫn "
    "được để trống thay vì suy diễn từ metric hành vi."
)

CONCLUSION = (
    "Bài báo trình bày SCT Camera và cập nhật đánh giá theo phiên bản hành vi đóng băng. Pilot "
    "một người gồm 20 clip (10 development, 10 test; 1.344,240 giây) cho thấy loitering và "
    "suspicious_behavior thành công trên test, trong khi intrusion còn sự kiện trùng, theft "
    "bỏ sót dương và báo sai âm, còn line_crossing có sự kiện sai/trùng. Véc-tơ Risk 22 chiều "
    "đã được audit đúng theo code, nhưng chưa có nhãn/model hợp lệ nên không báo cáo metric "
    "Risk. Các kết quả pilot này không xác lập khả năng khái quát theo người; bước tiếp theo "
    "là bổ sung người tham gia, khóa participant-disjoint split và sửa các lỗi hành vi đã quan "
    "sát trước khi thực hiện validation quy mô lớn."
)


BEHAVIOR_ROWS = [
    (
        "Intrusion",
        "Người đã xác nhận; ≥2 tâm liên tiếp; đoạn tâm cắt line hoặc line mới chạm bbox; chỉ IN đủ điều kiện alert.",
        "direction=IN → emit intrusion; không dùng polygon _analyze_zones.",
        "State side/touch theo camera-line-track; xóa khi track biến mất; không cooldown thời gian."
    ),
    (
        "Loitering",
        "Tâm người trong polygon; dwell mặc định 20 s; grace 3 s; policy vùng/danh tính có thể override.",
        "duration ≥ effective threshold → loitering/tier alert.",
        "Một alert mỗi tier/session; purge theo grace/session_gap."
    ),
    (
        "Suspicious behavior",
        "Confirmed stranger trong stranger_watch; ≥5 điểm; stationary ≤0,04 đường chéo hoặc path ≥0,18 và net ≤0,08; dwell 180 s.",
        "Đồng thời thỏa danh tính, hình học đáng ngờ và threshold thời gian.",
        "Một alert mỗi tier/session; purge theo grace/session_gap."
    ),
    (
        "Theft",
        "Stranger + vehicle trong asset_watch; khoảng cách ≤0,125 đường chéo; near ≥10 s; wrist conf ≥0,25 và cách bbox xe ≤0,055 tạo pose cue.",
        "Score ≥2/5; near bắt buộc và có ≥1 cue vehicle_started_moving/moving_same_direction/pose_push_contact; move ≥0,025, cosine ≥0,65.",
        "Một alert mỗi (camera, zone, vehicle class) đến khi reset; pair stale mặc định 30 s."
    ),
    (
        "Line crossing",
        "Người đã xác nhận; ≥2 tâm; cắt line hoặc mới chạm bbox; mapping forward/reverse tạo IN/OUT.",
        "IN/OUT đều tăng counter; production không emit line_crossing riêng; runner chỉ instrument counter.",
        "State side/touch theo track; xóa khi track biến mất."
    ),
]


FEATURE_ROWS = [
    (1, "duration — float(alert.duration), giây", 12, "displacement_ratio — max distance(first, point) / frame diagonal"),
    (2, "threshold_seconds — float(rule threshold), giây", 13, "speed_ratio — path_length_ratio / max(history_length−1, 1)"),
    (3, "duration_ratio — duration / max(threshold_seconds, 1)", 14, "has_vehicle_signal — OR của ba cue 15–17, {0,1}"),
    (4, "near_seconds — thời gian người gần tài sản, giây", 15, "vehicle_started_moving — cue dịch chuyển xe, {0,1}"),
    (5, "score — số cue luật; theft hiện tại 0…5", 16, "moving_same_direction — cue chuyển động tương quan, {0,1}"),
    (6, "score_ratio — score / max(score_threshold, 1)", 17, "pose_push_contact — cue tiếp xúc tư thế, {0,1}"),
    (7, "pacing_passes — số lần đổi phía khi đi qua lại", 18, "object_count — len(objects) trong khung hiện tại"),
    (8, "object_confidence — confidence; fallback identity_score", 19, "zone_configured — bool(camera_config.zones), {0,1}"),
    (9, "bbox_area_ratio — bbox area / frame area", 20, "is_person — class_name == person, {0,1}"),
    (10, "path_length_ratio — tổng bước Euclid / frame diagonal", 21, "is_vehicle — class_name ∈ VEHICLE_CLASSES, {0,1}"),
    (11, "net_distance_ratio — distance(first,last) / frame diagonal", 22, "is_asset — class_name ∈ ASSET_CLASSES, {0,1}"),
]


RESULT_ROWS = [
    ("Intrusion", "1/5/1/0", "1,000/0,000", "0,166667/1,000/0,285714", "0,457"),
    ("Loitering", "1/0/1/0", "1,000/0,000", "1,000/1,000/1,000", "0,725"),
    ("Suspicious behavior", "1/0/1/0", "1,000/0,000", "1,000/1,000/1,000", "0,186"),
    ("Theft", "0/1/0/1", "0,000/1,000", "0,000/0,000/0,000", "N/A"),
    ("Line crossing", "1/2/0/0", "1,000/1,000", "0,333333/1,000/0,500000", "4,084"),
]


def find_paragraph(document: Document, text: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}; found {len(matches)}")
    return matches[0]


def find_prefix(document: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting {prefix!r}; found {len(matches)}")
    return matches[0]


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph: Paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def insert_paragraph_after(paragraph: Paragraph, text: str, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    inserted.add_run(text)
    return inserted


def move_table_after(table: Table, paragraph: Paragraph) -> None:
    paragraph._p.addnext(table._tbl)


def set_cell_margins(cell, top: int = 45, start: int = 55, bottom: int = 45, end: int = 55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def format_table(table: Table, widths_cm: list[float], font_size: float) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.style = "TieuDeBang" if row_index == 0 else "NoiDungBang"
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)


def fill_cell(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold


def add_caption_after(document: Document, paragraph: Paragraph, text: str) -> Paragraph:
    return insert_paragraph_after(paragraph, text, "TenBang")


def build_behavior_table(document: Document, anchor: Paragraph) -> None:
    caption = add_caption_after(
        document,
        anchor,
        "Semantics và ngưỡng của năm hành vi trong phiên bản đóng băng.",
    )
    table = document.add_table(rows=1, cols=4)
    headers = ["Hành vi", "Bằng chứng/không gian/thời gian", "Trigger đóng băng", "Cooldown/Dedup"]
    for index, value in enumerate(headers):
        fill_cell(table.rows[0].cells[index], value, bold=True)
    for values in BEHAVIOR_ROWS:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            fill_cell(cells[index], value)
    format_table(table, [2.2, 5.1, 5.0, 3.5], 7.2)
    caption.paragraph_format.keep_with_next = True
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    move_table_after(table, caption)


def build_feature_table(document: Document, anchor: Paragraph) -> None:
    caption = add_caption_after(
        document,
        anchor,
        "Véc-tơ 22 đặc trưng theo đúng thứ tự và phép tính trong code runtime.",
    )
    table = document.add_table(rows=1, cols=4)
    headers = ["ID", "Feature / calculation", "ID", "Feature / calculation"]
    for index, value in enumerate(headers):
        fill_cell(table.rows[0].cells[index], value, bold=True)
    for left_id, left_text, right_id, right_text in FEATURE_ROWS:
        cells = table.add_row().cells
        for index, value in enumerate((str(left_id), left_text, str(right_id), right_text)):
            fill_cell(cells[index], value)
    format_table(table, [0.8, 7.1, 0.8, 7.1], 7.4)
    move_table_after(table, caption)


def replace_result_table(document: Document, old_table: Table, caption: Paragraph) -> None:
    old_table._tbl.getparent().remove(old_table._tbl)
    table = document.add_table(rows=1, cols=5)
    headers = ["Hành vi", "TP/FP/TN/FN", "DR/FAR", "Precision/Recall/F1", "Latency median (s)"]
    for index, value in enumerate(headers):
        fill_cell(table.rows[0].cells[index], value, bold=True)
    for values in RESULT_ROWS:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            fill_cell(cells[index], value)
    format_table(table, [3.4, 2.7, 2.4, 4.6, 2.7], 7.8)
    move_table_after(table, caption)


def replace_picture(paragraph: Paragraph, figure_path: Path) -> None:
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(figure_path), width=Cm(15.8))


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: update_fdse_paper_frozen_pilot.py SOURCE.docx FIGURE.png OUTPUT.docx")
    source = Path(sys.argv[1]).resolve()
    figure = Path(sys.argv[2]).resolve()
    output = Path(sys.argv[3]).resolve()
    if source == output:
        raise RuntimeError("Output must differ from source")
    if not source.exists() or not figure.exists():
        raise FileNotFoundError(source if not source.exists() else figure)

    document = Document(source)
    if len(document.tables) != 3:
        raise RuntimeError(f"Expected 3 source tables, found {len(document.tables)}")

    set_text(find_paragraph(document, "Tóm tắt"), "Tóm tắt")
    set_text(find_prefix(document, "Giám sát video quy mô nhỏ cần vượt"), ABSTRACT)
    set_text(find_prefix(document, "(3) Bộ máy hành vi ưu tiên"), CONTRIBUTION_3)
    set_text(find_prefix(document, "(4) Cổng rủi ro hồi quy logistic"), CONTRIBUTION_4)
    set_text(find_prefix(document, "(5) Đánh giá tái lập toàn bộ"), CONTRIBUTION_5)
    set_text(find_prefix(document, "SCT Camera kết hợp các luật theo camera"), RELATED_RISK)
    set_text(find_prefix(document, "Hình 1 tóm tắt sáu giai đoạn"), ARCHITECTURE)

    behavior_anchor = find_prefix(document, "Sự kiện cắt đường phát khi")
    set_text(behavior_anchor, BEHAVIOR_INTRO)
    build_behavior_table(document, behavior_anchor)

    remove_paragraph(find_paragraph(document, "Chấm điểm rủi ro hành vi bằng hồi quy logistic trên véc-tơ 22 đặc trưng."))
    # The first empty drawing paragraph after the risk heading is the obsolete Risk figure.
    risk_heading = find_paragraph(document, "Học hành vi và cổng rủi ro")
    candidate = risk_heading._p.getnext()
    if candidate is None or not candidate.xpath(".//w:drawing"):
        raise RuntimeError("Risk figure paragraph not found after risk heading")
    remove_paragraph(Paragraph(candidate, risk_heading._parent))

    risk_intro = find_prefix(document, "Mỗi ứng viên được biểu diễn")
    set_text(risk_intro, RISK_INTRO)
    build_feature_table(document, risk_intro)
    for exact in ("\t\t(5)", "\t\t(6)", "\t\t(7)"):
        remove_paragraph(find_paragraph(document, exact))
    remove_paragraph(find_prefix(document, "scripts/train_behavior_classifier.py"))
    remove_paragraph(find_prefix(document, "trong đó mỗi bước cập nhật"))
    set_text(find_prefix(document, "Mô hình gồm tham số chuẩn hóa"), RISK_LIMITATION)

    set_text(find_paragraph(document, "Học hành vi và cổng rủi ro"), "Đặc trưng hành vi và cổng rủi ro")
    alert_caption = find_prefix(document, "AlertManager hướng sự kiện qua")
    alert_caption.paragraph_format.keep_with_next = False
    alert_caption.paragraph_format.space_before = Pt(0)
    alert_caption.paragraph_format.space_after = Pt(0)
    alert_picture_element = alert_caption._p.getprevious()
    if alert_picture_element is None or not alert_picture_element.xpath(".//w:drawing"):
        raise RuntimeError("AlertManager figure paragraph not found before caption")
    alert_picture = Paragraph(alert_picture_element, alert_caption._parent)
    alert_picture.paragraph_format.keep_with_next = True
    alert_picture.paragraph_format.space_before = Pt(0)
    alert_picture.paragraph_format.space_after = Pt(0)
    if len(document.inline_shapes) != 6:
        raise RuntimeError(f"Expected 6 inline shapes before resize, found {len(document.inline_shapes)}")
    document.inline_shapes[2].width = Cm(5.8)
    document.inline_shapes[2].height = Cm(2.5)

    set_text(find_paragraph(document, "Môi trường"), "Môi trường và thiết lập thí nghiệm")
    environment = find_prefix(document, "Phép đo chạy trên Windows 10")
    set_text(environment, ENVIRONMENT)
    insert_paragraph_after(environment, PILOT_SETUP, "NoiDung")

    set_text(find_prefix(document, "Các luồng đẩy cảnh báo qua"), ALERT_DISTRIBUTION)
    set_text(find_prefix(document, "RSS đạt cực đại"), find_prefix(document, "RSS đạt cực đại").text.replace("Hình 5d", "Hình 4d"))
    set_text(find_prefix(document, "Hình 5 cho thấy"), find_prefix(document, "Hình 5 cho thấy").text.replace("Hình 5", "Hình 4", 1))
    set_text(find_prefix(document, "Bảng 2 cho thấy"), find_prefix(document, "Bảng 2 cho thấy").text.replace("Bảng 2", "Bảng 4", 1))
    set_text(find_prefix(document, "Độ trễ danh tính trong Bảng 2"), find_prefix(document, "Độ trễ danh tính trong Bảng 2").text.replace("Bảng 2", "Bảng 4", 1))
    set_text(find_prefix(document, "pytest thực thi 135 ca"), REGRESSION_RESULTS)

    set_text(find_paragraph(document, "Kết quả thực nghiệm ở mức hành vi"), "Pilot hành vi trên tập test đóng băng")
    set_text(find_prefix(document, "Thí nghiệm hành vi xử lý trực tuyến"), RESULT_SETUP)
    set_text(find_prefix(document, "Mỗi cặp người lạ đã xác nhận"), FIGURE_EXPLANATION)

    result_picture = find_paragraph(document, "  ")
    replace_picture(result_picture, figure)
    set_text(find_prefix(document, "Diễn tiến điểm trộm cắp"), "Pilot behavior evaluation outcomes on the frozen test split.")
    set_text(find_prefix(document, "Hình 6(a):"), RESULT_HIGHLIGHTS)
    set_text(find_prefix(document, "Hình 6(b):"), "Giá trị DR/FAR và Precision/Recall/F1 trong bảng lần lượt dùng đơn vị sự kiện/clip âm như định nghĩa ở trên; N/A nghĩa là không có TP để tính latency.")
    result_caption = find_prefix(document, "Các chỉ số quan sát trong Hình 6")
    set_text(result_caption, "Kết quả từng hành vi trên tập test đóng băng (mỗi hành vi: một clip dương, một clip âm).")
    old_result_table = document.tables[-1]
    replace_result_table(document, old_result_table, result_caption)
    set_text(find_prefix(document, "Hình 6(a-b) cho thấy"), "Không báo cáo điểm tổng hợp cho pilot; các giá trị phải được đọc theo từng hành vi và cùng kết quả âm trong Hình 5.")

    set_text(find_prefix(document, "Ba hạn chế chính vẫn tồn tại"), DISCUSSION_1)
    set_text(find_prefix(document, "Khóa suy luận YOLO dùng chung"), DISCUSSION_2)
    set_text(find_prefix(document, "Bài báo trình bày SCT Camera, quy trình sáu giai đoạn"), CONCLUSION)

    references_started = False
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "TÀI LIỆU THAM KHẢO":
            references_started = True
            continue
        if references_started and paragraph.style.name == "TaiLieuThamKhao":
            paragraph.paragraph_format.space_after = Pt(0)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    check = Document(output)
    required = [
        "1.344,240",
        "priority-pilot-v2-single-participant-2026-08-18",
        "Pilot behavior evaluation outcomes on the frozen test split.",
        "Risk Accuracy, Precision, Recall, F1, ROC-AUC",
        "participant-level generalization",
    ]
    corpus = "\n".join(paragraph.text for paragraph in check.paragraphs)
    corpus += "\n" + "\n".join(cell.text for table in check.tables for row in table.rows for cell in row.cells)
    missing = [value for value in required if value not in corpus]
    if missing:
        raise RuntimeError(f"Missing required content after save: {missing}")
    if len(check.tables) != 5:
        raise RuntimeError(f"Expected 5 output tables, found {len(check.tables)}")
    print(f"output={output}")
    print(f"paragraphs={len(check.paragraphs)} tables={len(check.tables)} inline_shapes={len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
