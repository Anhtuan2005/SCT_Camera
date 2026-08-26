from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document


FEATURES = [
    "duration", "threshold_seconds", "duration_ratio", "near_seconds", "score",
    "score_ratio", "pacing_passes", "object_confidence", "bbox_area_ratio",
    "path_length_ratio", "net_distance_ratio", "displacement_ratio", "speed_ratio",
    "has_vehicle_signal", "vehicle_started_moving", "moving_same_direction",
    "pose_push_contact", "object_count", "zone_configured", "is_person",
    "is_vehicle", "is_asset",
]

REQUIRED = [
    "20 clip",
    "10 clip phát triển",
    "10 clip test đóng băng",
    "1.344,240",
    "666,296",
    "677,944",
    "priority-pilot-v2-single-participant-2026-08-18",
    "Pilot behavior evaluation outcomes on the frozen test split.",
    "năm sự kiện IN trùng",
    "Theft bỏ sót clip dương",
    "participant-level generalization",
    "Risk Accuracy, Precision, Recall, F1, ROC-AUC",
    "cơ chế loại bỏ sự kiện trùng lặp theo track ID chưa ổn định",
    "Đánh giá hồi quy lịch sử trong benchmark cũ",
    "206/206 ca, 0 thất bại, trong 8,16 giây",
]

FORBIDDEN = [
    "Thí nghiệm hành vi xử lý trực tuyến tương tác người-xe máy",
    "Diễn tiến điểm trộm cắp trên luồng trực tiếp",
    "Hình 6(a)",
    "Hình 6(b)",
    "xâm nhập cảnh báo khi người vào vùng hạn chế",
    "mô hình hồi quy logistic 22 đặc trưng học từ sự kiện do người vận hành gán nhãn",
    "NEEDS MANUAL INPUT",
    "track ID chưa khử trùng ổn định",
]


def main(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        corrupt = archive.testzip()
        if corrupt:
            raise RuntimeError(f"Corrupt package part: {corrupt}")

    document = Document(path)
    corpus = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_texts = [
        "\n".join(cell.text for row in table.rows for cell in row.cells)
        for table in document.tables
    ]
    corpus += "\n" + "\n".join(table_texts)

    if len(document.tables) != 5:
        raise RuntimeError(f"Expected 5 tables, found {len(document.tables)}")
    if len(document.inline_shapes) != 5:
        raise RuntimeError(f"Expected 5 figures, found {len(document.inline_shapes)}")
    missing = [value for value in REQUIRED if value not in corpus]
    if missing:
        raise RuntimeError(f"Missing required content: {missing}")
    forbidden = [value for value in FORBIDDEN if value in corpus]
    if forbidden:
        raise RuntimeError(f"Obsolete content remains: {forbidden}")
    missing_features = [feature for feature in FEATURES if feature not in table_texts[1]]
    if missing_features:
        raise RuntimeError(f"Missing audited features: {missing_features}")

    result_table = document.tables[4]
    expected_rows = {
        "Intrusion": ["1/5/1/0", "1,000/0,000", "0,166667/1,000/0,285714", "0,457"],
        "Loitering": ["1/0/1/0", "1,000/0,000", "1,000/1,000/1,000", "0,725"],
        "Suspicious behavior": ["1/0/1/0", "1,000/0,000", "1,000/1,000/1,000", "0,186"],
        "Theft": ["0/1/0/1", "0,000/1,000", "0,000/0,000/0,000", "N/A"],
        "Line crossing": ["1/2/0/0", "1,000/1,000", "0,333333/1,000/0,500000", "4,084"],
    }
    actual_rows = {
        row.cells[0].text: [cell.text for cell in row.cells[1:]]
        for row in result_table.rows[1:]
    }
    if actual_rows != expected_rows:
        raise RuntimeError(f"Frozen result table mismatch: {actual_rows}")

    print(
        f"verified={path} paragraphs={len(document.paragraphs)} "
        f"tables={len(document.tables)} figures={len(document.inline_shapes)} features=22"
    )


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: verify_fdse_paper_frozen_pilot.py PAPER.docx")
    main(Path(sys.argv[1]).resolve())
