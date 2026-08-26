from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"E:\SCT_Camera\NOI_DUNG_BO_SUNG_BAO_CAO.docx")
OUT_DIR = Path(r"E:\SCT_Camera\.codex_addendum_baocao_20260821_01\debug")


def make_variant(limit: int) -> Path:
    source = Document(SOURCE)
    target = Document()
    target.element.body.clear_content()
    body_children = [
        child for child in source.element.body if child.tag != qn("w:sectPr")
    ]
    for child in body_children[:limit]:
        target.element.body.append(deepcopy(child))
    target.element.body.append(deepcopy(source.element.body.sectPr))
    path = OUT_DIR / f"variant-{limit:03d}.docx"
    target.save(path)
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source = Document(SOURCE)
    children = [child for child in source.element.body if child.tag != qn("w:sectPr")]
    print(f"Body elements: {len(children)}")
    for limit in [1, 20, 40, 60, 80, 100, len(children)]:
        path = make_variant(limit)
        print(path)


if __name__ == "__main__":
    main()
