from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"E:\SCT_Camera\NOI_DUNG_BO_SUNG_BAO_CAO.docx")
FONT = "Times New Roman"
USABLE_CM = 15.5


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_style_font(style, size, bold=None, italic=None):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    doc.core_properties.title = "Nội dung bổ sung đề xuất cho báo cáo khóa luận"
    doc.core_properties.subject = "Ablation, phân tích lỗi, tái lập, quyền riêng tư và tính hợp lệ"
    doc.core_properties.author = "Nhóm sinh viên thực hiện"


def shade_paragraph(paragraph, fill="FFF2CC"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B7A04A")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_body(doc, text, *, first_indent=True):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.0 if first_indent else 0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run(text), 13)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(text), 13)
    return p


def add_placeholder(paragraph, text):
    run = paragraph.add_run(f"[[BỔ SUNG: {text}]]")
    set_run_font(run, 13, bold=True)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(text), 15 if level == 1 else 14, bold=True, color=(0, 0, 0))
    return p


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_cm):
    assert abs(sum(widths_cm) - USABLE_CM) < 0.02
    widths_twips = [round(width / 2.54 * 1440) for width in widths_cm]
    total = sum(widths_twips)
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[index])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_twips[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc, caption, headers, rows, widths_cm, *, placeholder_columns=()):
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.keep_together = True
    set_run_font(cap.add_run(caption), 12, italic=True, color=(0, 0, 0))
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        shade_cell(cell, "D9EAF7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), 11.5, bold=True)
    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == len(values) - 1 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, 11.5)
            if index in placeholder_columns or str(value).startswith("[[BỔ SUNG"):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                run.bold = True
    set_table_geometry(table, widths_cm)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = Cm(0)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(12)
    set_run_font(title.add_run("NỘI DUNG BỔ SUNG ĐỀ XUẤT CHO BÁO CÁO"), 15, bold=True)

    note = doc.add_paragraph(style="Normal")
    note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.line_spacing = 1.15
    note.paragraph_format.space_after = Pt(10)
    shade_paragraph(note)
    set_run_font(note.add_run("Ghi chú sử dụng: "), 12, bold=True, italic=True)
    set_run_font(
        note.add_run(
            "Tài liệu này chỉ chứa các phần bổ sung để chèn vào báo cáo hiện có. "
            "Các tiêu đề dùng ký hiệu 4.x/5.x để tránh làm sai số mục trong bản chính. "
            "Mọi nội dung tô vàng phải được thay bằng số liệu đo thực tế trước khi nộp; không được suy đoán hoặc tự tạo kết quả."
        ),
        12,
        italic=True,
    )

    add_heading(doc, "CÁC MỤC BỔ SUNG CHO CHƯƠNG 4", 1)
    add_heading(doc, "4.x Đánh giá đóng góp của từng thành phần", 2)
    add_body(
        doc,
        "Các kết quả hiện có cho thấy hệ thống hoàn chỉnh có thể vận hành với nhiều camera, thực hiện phát hiện, theo dõi, nhận diện danh tính và phân tích hành vi trong cùng một pipeline. Tuy nhiên, phép đo đầu cuối chưa cho biết mức cải thiện đến từ ByteTrack gốc hay từ từng cơ chế bổ sung. Vì vậy, thí nghiệm loại bỏ thành phần (ablation study) được đề xuất nhằm tách riêng ảnh hưởng của bộ nhớ track bị mất tạm thời, cơ chế bù chuyển động camera và bước loại bỏ khung bao trùng lặp.",
    )
    add_body(
        doc,
        "Nguyên tắc của thí nghiệm là chỉ thay đổi một thành phần tại một thời điểm, trong khi giữ nguyên video đầu vào, checkpoint YOLO11s, cấu hình ByteTrack, ngưỡng confidence, kích thước ảnh suy luận, phần cứng và cách ghi metric. Cách tổ chức này giúp giảm khả năng quy kết nhầm một thay đổi về kết quả cho các yếu tố ngoài thành phần đang được đánh giá.",
    )

    add_heading(doc, "4.x.1 Mục tiêu và giả thuyết đánh giá", 3)
    add_body(doc, "Thí nghiệm tập trung kiểm tra ba giả thuyết kỹ thuật sau:")
    add_bullet(doc, "Bộ nhớ track bị mất tạm thời giúp giảm số quỹ đạo bị đứt và số lần cấp lại ID khi đối tượng bị che khuất trong thời gian ngắn.")
    add_bullet(doc, "Cơ chế bù chuyển động camera giúp giảm sai lệch dự đoán vị trí và giảm ID switch trong các đoạn video có rung hoặc xoay nhẹ.")
    add_bullet(doc, "Bước lọc khung bao trùng lặp giúp giảm số phát hiện dư thừa mà không làm tăng đáng kể số đối tượng bị bỏ sót.")
    add_body(
        doc,
        "Ngoài chất lượng theo dõi, thí nghiệm cũng phải kiểm tra chi phí tính toán của từng cơ chế. Một cải tiến chỉ được xem là hữu ích khi mức tăng chất lượng đủ lớn so với phần độ trễ, mức sử dụng CPU/GPU hoặc bộ nhớ mà nó bổ sung.",
    )

    add_heading(doc, "4.x.2 Cấu hình thí nghiệm", 3)
    add_table(
        doc,
        "Bảng 4-x: Cấu hình thí nghiệm loại bỏ thành phần",
        ["Mã", "ByteTrack", "Bộ nhớ track", "CMC", "Lọc box trùng", "Mục đích"],
        [
            ["A0", "Có", "Không", "Không", "Không", "Baseline để xác định chất lượng và chi phí của ByteTrack gốc."],
            ["A1", "Có", "Có", "Không", "Không", "Đo riêng tác dụng của bộ nhớ track bị mất tạm thời."],
            ["A2", "Có", "Không", "Có", "Không", "Đo riêng tác dụng của bù chuyển động camera."],
            ["A3", "Có", "Không", "Không", "Có", "Đo riêng tác dụng của bước lọc khung bao trùng lặp."],
            ["A4", "Có", "Có", "Có", "Có", "Cấu hình hoàn chỉnh đang được sử dụng trong hệ thống."],
        ],
        [1.1, 2.2, 2.3, 1.5, 2.1, 6.3],
    )
    add_body(
        doc,
        "Tập video dùng cho ablation cần có đủ ba nhóm tình huống: đối tượng bị che khuất, camera có rung nhẹ và cảnh xuất hiện nhiều khung bao chồng lặp. Nếu dùng lại frozen test hiện có, cần bổ sung clip chuyên biệt cho chất lượng tracking vì frozen test hành vi chỉ phản ánh kết quả sự kiện cuối cùng và không đủ để quan sát mọi ID switch hoặc quỹ đạo bị đứt.",
    )
    p = add_body(doc, "Phạm vi dữ liệu dự kiến cho phép đo: ")
    add_placeholder(p, "số video, tổng thời lượng, độ phân giải, FPS và số track ground truth")

    add_heading(doc, "4.x.3 Chỉ số đánh giá", 3)
    add_table(
        doc,
        "Bảng 4-x: Các chỉ số dùng trong thí nghiệm ablation",
        ["Chỉ số", "Đơn vị", "Ý nghĩa trong thí nghiệm"],
        [
            ["IDF1", "%", "Mức nhất quán danh tính của đối tượng trong toàn bộ chuỗi video."],
            ["ID switch", "Lần", "Số lần một đối tượng thật bị đổi sang ID khác hoặc bị gán nhầm ID."],
            ["Fragmentation", "Lần", "Số lần một quỹ đạo liên tục bị chia thành nhiều đoạn track."],
            ["Duplicate boxes", "Box", "Số khung bao dư thừa cho cùng một đối tượng sau hậu xử lý."],
            ["AI FPS", "Khung hình/giây", "Thông lượng xử lý của pipeline với từng cấu hình."],
            ["Độ trễ p50/p95", "ms", "Độ trễ điển hình và độ trễ ở nhóm khung hình xử lý chậm."],
            ["VRAM cực đại", "MB", "Chi phí bộ nhớ GPU lớn nhất trong cửa sổ đo."],
        ],
        [3.0, 2.4, 10.1],
    )
    add_body(
        doc,
        "IDF1, ID switch và fragmentation phải được tính từ cùng một tập ground truth. Duplicate boxes được tính sau bước gán track để phản ánh đúng vai trò của cơ chế lọc trùng. AI FPS và độ trễ phải được đo sau giai đoạn warm-up, tách khỏi thời gian khởi tạo mô hình và không tính thời gian người vận hành chuẩn bị camera.",
    )

    add_heading(doc, "4.x.4 Giao thức thực hiện và kiểm soát sai lệch", 3)
    add_body(
        doc,
        "Mỗi cấu hình được chạy trên cùng thứ tự video và cùng cấu hình phần cứng. Cache, tác vụ nền, ghi hình và kênh gửi thông báo phải được giữ cố định giữa các lần chạy. Nếu có thành phần bất đồng bộ, cần đợi hàng đợi xử lý trở về trạng thái rỗng trước khi bắt đầu phép đo tiếp theo.",
    )
    add_bullet(doc, "Khóa checkpoint và ghi lại mã băm của file mô hình.")
    add_bullet(doc, "Khóa phiên bản mã nguồn, phiên bản thư viện và file cấu hình tracker.")
    add_bullet(doc, "Dùng cùng một vùng ROI, vạch giám sát và ngưỡng confidence cho mọi cấu hình.")
    add_bullet(doc, "Chạy một lượt warm-up trước khi thu thập thời gian xử lý.")
    add_bullet(doc, "Lặp lại mỗi cấu hình đủ số lần và báo cáo cả trung bình lẫn độ phân tán.")
    p = add_body(doc, "Số lần lặp cho mỗi cấu hình: ")
    add_placeholder(p, "N lần; khuyến nghị tối thiểu 3 lần nếu điều kiện cho phép")

    add_heading(doc, "4.x.5 Kết quả thí nghiệm", 3)
    add_table(
        doc,
        "Bảng 4-x: Kết quả chất lượng tracking của các cấu hình ablation",
        ["Cấu hình", "IDF1 (%)", "ID switch", "Fragmentation", "Box trùng"],
        [
            ["A0", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A1", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A2", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A3", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A4", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
        ],
        [2.5, 3.0, 3.0, 3.6, 3.4],
        placeholder_columns=(1, 2, 3, 4),
    )
    add_table(
        doc,
        "Bảng 4-x: Chi phí xử lý của các cấu hình ablation",
        ["Cấu hình", "AI FPS", "p50 (ms)", "p95 (ms)", "VRAM (MB)"],
        [
            ["A0", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A1", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A2", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A3", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
            ["A4", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]", "[[BỔ SUNG]]"],
        ],
        [2.5, 3.0, 3.0, 3.0, 4.0],
        placeholder_columns=(1, 2, 3, 4),
    )
    p = add_body(doc, "Nhận xét chính sau khi có số liệu: ")
    add_placeholder(p, "nêu cấu hình cải thiện IDF1/ID switch nhiều nhất và chi phí độ trễ tương ứng")
    p = add_body(doc, "Kết luận của ablation: ")
    add_placeholder(p, "chỉ kết luận thành phần có hiệu quả khi chênh lệch lặp lại ổn định qua các lần chạy")

    page_break(doc)
    add_heading(doc, "4.x Phân tích lỗi định tính", 2)
    add_body(
        doc,
        "Các metric tổng hợp cho biết hệ thống đúng hoặc sai bao nhiêu lần nhưng không giải thích đầy đủ vì sao lỗi xảy ra. Phân tích lỗi định tính bổ sung góc nhìn theo chuỗi xử lý: lỗi ở tầng phát hiện có thể lan sang tracking, nhận diện danh tính, phân tích hành vi và cuối cùng biểu hiện dưới dạng một cảnh báo sai. Vì vậy, đơn vị phân tích cần bao gồm khung hình, track và sự kiện thay vì chỉ xem từng ảnh riêng lẻ.",
    )

    add_heading(doc, "4.x.1 Quy trình ghi nhận ca lỗi", 3)
    add_body(doc, "Mỗi ca lỗi được ghi nhận theo năm bước thống nhất:")
    add_bullet(doc, "Xác định thời điểm bắt đầu và kết thúc của hiện tượng trong video.")
    add_bullet(doc, "Đối chiếu kết quả hệ thống với ground truth ở mức box, track và sự kiện.")
    add_bullet(doc, "Xác định tầng đầu tiên tạo ra sai lệch, tránh quy toàn bộ lỗi cho module cảnh báo cuối.")
    add_bullet(doc, "Lưu lại ảnh trước, trong và sau thời điểm lỗi để quan sát diễn biến theo thời gian.")
    add_bullet(doc, "Ghi biện pháp giảm lỗi và ảnh hưởng phụ có thể phát sinh khi thay đổi ngưỡng hoặc logic.")

    add_heading(doc, "4.x.2 Phân loại lỗi", 3)
    add_table(
        doc,
        "Bảng 4-x: Phân loại lỗi theo tầng xử lý",
        ["Nhóm lỗi", "Biểu hiện", "Nguyên nhân có thể", "Metric liên quan"],
        [
            ["Phát hiện", "Bỏ sót hoặc nhận nhầm đối tượng.", "Vật thể nhỏ, thiếu sáng, mờ chuyển động, che khuất.", "FN, FP, recall"],
            ["Theo dõi", "Đổi ID hoặc gãy quỹ đạo.", "Che khuất, camera rung, hai người giao nhau.", "IDF1, ID switch, fragmentation"],
            ["Danh tính", "Known/stranger không ổn định.", "Khuôn mặt nhỏ, góc nghiêng, ít ảnh enrollment.", "FAR, FRR, score cosine"],
            ["Hành vi", "Thiếu sự kiện hoặc cảnh báo lặp.", "Ngưỡng thời gian, ROI, ghép sự kiện chưa phù hợp.", "TP, FP, FN, FAR"],
            ["Thông báo", "Gửi chậm, gửi trùng hoặc mất thông báo.", "I/O mạng, retry, hàng đợi chưa được giới hạn.", "Độ trễ, số lần gửi trùng"],
        ],
        [2.3, 4.0, 5.7, 3.5],
    )

    add_heading(doc, "4.x.3 Che khuất, giao cắt quỹ đạo và ID switch", 3)
    add_body(
        doc,
        "Khi hai người đi sát hoặc cắt ngang nhau, khung bao có thể chồng lấn và độ tương đồng IoU giữa track dự đoán với detection mới trở nên gần nhau. Nếu một người bị che khuất trong vài khung hình, tracker có thể gán detection vừa xuất hiện cho track còn lại hoặc tạo ID mới. Ca lỗi cần được phân tích trên một cửa sổ thời gian đủ dài để xác định ID switch là lỗi gán nhất thời hay một quỹ đạo đã bị gãy hoàn toàn.",
    )
    add_body(
        doc,
        "Bộ nhớ track bị mất tạm thời chỉ có tác dụng trong khoảng ân hạn. Nếu tăng khoảng này quá lớn, hệ thống có thể ghép nhầm một người mới vào track cũ; nếu đặt quá nhỏ, đối tượng vừa bị che khuất sẽ bị cấp lại ID. Vì vậy, khi mô tả một ca lỗi phải ghi cả thời gian che khuất, khoảng cách vị trí, biến đổi kích thước box và giá trị ngưỡng đang sử dụng.",
    )

    add_heading(doc, "4.x.4 Vật thể nhỏ, ánh sáng yếu và mờ chuyển động", 3)
    add_body(
        doc,
        "Vật thể ở xa chiếm ít pixel nên đặc trưng dùng cho phát hiện và nhận diện đều suy giảm. Trong điều kiện thiếu sáng, nhiễu ảnh làm biên vật thể kém rõ; khi camera hoặc đối tượng chuyển động nhanh, box giữa hai khung liên tiếp có thể lệch đủ lớn để làm giảm IoU. Ba hiện tượng này thường xuất hiện đồng thời nên cần lưu cả kích thước box, confidence và mức thay đổi vị trí theo thời gian.",
    )
    add_body(
        doc,
        "Biện pháp khắc phục phải được đánh giá theo đánh đổi. Giảm ngưỡng confidence có thể tăng recall nhưng cũng làm tăng số box nền; tăng độ phân giải suy luận có thể giúp vật thể nhỏ nhưng làm giảm FPS; tăng độ nhạy của tracker có thể giảm fragmentation nhưng làm tăng nguy cơ ghép nhầm. Báo cáo nên nêu rõ đánh đổi thay vì chỉ mô tả một tham số là tốt hoặc xấu.",
    )

    add_heading(doc, "4.x.5 Lỗi nhận diện danh tính", 3)
    add_body(
        doc,
        "Đối với nhận diện khuôn mặt, cần tách ba tình huống: không phát hiện được khuôn mặt, phát hiện được nhưng embedding không đủ chất lượng, và embedding có score gần ngưỡng quyết định. Việc chỉ xem nhãn known/stranger ở đầu ra sẽ che khuất nguyên nhân thật. Mỗi ca lỗi nên ghi detection score của khuôn mặt, cosine similarity lớn nhất, danh tính ứng viên và số lần xác nhận đã tích lũy trên track.",
    )
    add_body(
        doc,
        "Bộ nhớ danh tính ngắn hạn giúp giảm hiện tượng nhấp nháy nhãn khi tracker cấp lại ID, nhưng cũng có thể truyền nhãn sai nếu hai người xuất hiện gần cùng vị trí trong khoảng thời gian ngắn. Vì vậy, các ca liên quan đến kế thừa danh tính cần được đánh dấu riêng và kiểm tra đồng thời khoảng cách không gian, thời gian và tỉ lệ khung bao.",
    )

    add_heading(doc, "4.x.6 Lỗi sự kiện và cảnh báo", 3)
    add_body(
        doc,
        "Một cảnh báo sai không nhất thiết bắt nguồn từ mô hình AI. Ví dụ, detection và track có thể đúng nhưng ROI được vẽ lệch; sự kiện có thể đúng nhưng bộ ghép sự kiện phát lặp; hoặc pipeline có thể tạo một cảnh báo hợp lệ nhưng kênh mạng gửi lại do retry. Do đó, log phân tích cần lưu mã camera, ID track, loại sự kiện, thời điểm, khóa chống trùng và trạng thái gửi của từng kênh.",
    )
    add_table(
        doc,
        "Bảng 4-x: Mẫu tổng hợp các ca lỗi đại diện",
        ["Mã", "Điều kiện", "Quan sát", "Nguyên nhân và hướng xử lý"],
        [
            ["E01", "Che khuất/giao cắt", "[[BỔ SUNG ẢNH VÀ MÔ TẢ]]", "[[BỔ SUNG NGUYÊN NHÂN, BIỆN PHÁP]]"],
            ["E02", "Thiếu sáng", "[[BỔ SUNG ẢNH VÀ MÔ TẢ]]", "[[BỔ SUNG NGUYÊN NHÂN, BIỆN PHÁP]]"],
            ["E03", "Camera rung", "[[BỔ SUNG ẢNH VÀ MÔ TẢ]]", "[[BỔ SUNG NGUYÊN NHÂN, BIỆN PHÁP]]"],
            ["E04", "Khuôn mặt nhỏ/nghiêng", "[[BỔ SUNG ẢNH VÀ MÔ TẢ]]", "[[BỔ SUNG NGUYÊN NHÂN, BIỆN PHÁP]]"],
            ["E05", "Cảnh báo lặp", "[[BỔ SUNG ẢNH VÀ MÔ TẢ]]", "[[BỔ SUNG NGUYÊN NHÂN, BIỆN PHÁP]]"],
        ],
        [1.2, 3.3, 5.0, 6.0],
        placeholder_columns=(2, 3),
    )

    page_break(doc)
    add_heading(doc, "4.x Độ tin cậy thống kê và khả năng tái lập", 2)
    add_body(
        doc,
        "Báo cáo kết quả thực nghiệm cần phân biệt phép đo hiệu năng lặp lại nhiều lần với pilot có sample size nhỏ. Đối với FPS và độ trễ, mỗi cấu hình có thể được chạy lặp để báo cáo trung bình, độ lệch chuẩn và khoảng biến thiên. Đối với frozen test hành vi hiện chỉ có ít sự kiện dương tính, việc tính một điểm tổng hợp hoặc khoảng tin cậy có thể tạo cảm giác chắc chắn lớn hơn dữ liệu thực tế; vì vậy kết quả nên tiếp tục được trình bày theo từng hành vi và kèm số đếm TP, FP, FN.",
    )
    add_body(
        doc,
        "Khả năng tái lập phụ thuộc vào việc lưu đủ thông tin để một lần chạy có thể được tái tạo. Ngoài tên mô hình và phần cứng, cần ghi phiên bản mã nguồn, mã băm checkpoint, phiên bản thư viện, file cấu hình, ngưỡng xử lý, định danh dữ liệu và thời gian đo. Nếu một trong các yếu tố này thay đổi, kết quả mới phải được xem là một cấu hình thí nghiệm khác.",
    )

    add_heading(doc, "4.x.1 Báo cáo độ phân tán", 3)
    add_body(
        doc,
        "Với metric thời gian, giá trị trung bình mô tả xu hướng chung nhưng không phản ánh các khung hình chậm bất thường. Vì vậy, báo cáo nên kèm p50, p95 và p99. Khi có từ ba lần chạy trở lên, có thể báo cáo trung bình ± độ lệch chuẩn cho FPS và độ trễ; đồng thời giữ số liệu từng lần chạy trong phụ lục để tránh che khuất một lần chạy bất thường.",
    )
    p = add_body(doc, "Quy ước báo cáo số liệu lặp: ")
    add_placeholder(p, "N lần chạy, thời lượng mỗi lần, cách tính warm-up và cửa sổ đo")

    add_heading(doc, "4.x.2 Thông tin tối thiểu để tái lập", 3)
    add_table(
        doc,
        "Bảng 4-x: Thông tin cần khóa cho mỗi phép đo",
        ["Nhóm", "Thông tin cần ghi", "Giá trị của lần đo"],
        [
            ["Mã nguồn", "Commit/tag và trạng thái working tree", "[[BỔ SUNG]]"],
            ["Mô hình", "Tên checkpoint và SHA-256", "[[BỔ SUNG]]"],
            ["Phần mềm", "Python, CUDA, cuDNN, PyTorch, Ultralytics", "[[BỔ SUNG]]"],
            ["Phần cứng", "CPU, GPU, RAM, hệ điều hành", "[[BỔ SUNG]]"],
            ["Dữ liệu", "Tên tập, mã băm manifest, số clip, tổng thời lượng", "[[BỔ SUNG]]"],
            ["Cấu hình", "Confidence, IoU, tracker, ROI, thời gian ân hạn", "[[BỔ SUNG]]"],
            ["Đo lường", "Warm-up, số lần lặp, cửa sổ đo, tác vụ nền", "[[BỔ SUNG]]"],
        ],
        [2.6, 7.2, 5.7],
        placeholder_columns=(2,),
    )

    add_heading(doc, "4.x.3 Phạm vi diễn giải kết quả", 3)
    add_body(
        doc,
        "Kết quả pilot hành vi hiện phản ánh khả năng hoạt động trên một tập đóng băng và một phạm vi điều kiện cụ thể, chưa phải ước lượng hiệu năng trên quần thể người dùng hoặc địa điểm mới. Tương tự, pilot khuôn mặt với hai danh tính chỉ phù hợp để kiểm tra pipeline và lựa chọn ngưỡng sơ bộ. Các kết luận trong báo cáo cần dùng các cụm từ như 'trong phạm vi pilot' hoặc 'trên tập frozen test hiện tại' thay cho các phát biểu tổng quát về độ chính xác ngoài thực tế.",
    )
    add_body(
        doc,
        "Đối với phép đo bốn camera, kết luận chỉ nên áp dụng cho cấu hình phần cứng, độ phân giải và tải xử lý đã nêu. Việc suy rộng sang số lượng camera lớn hơn cần thêm phép đo scaling hoặc một mô hình tải được kiểm chứng bằng dữ liệu thực tế.",
    )

    page_break(doc)
    add_heading(doc, "CÁC MỤC BỔ SUNG CHO CHƯƠNG 5", 1)
    add_heading(doc, "5.x Quyền riêng tư, bảo mật và đạo đức triển khai", 2)
    add_body(
        doc,
        "Hệ thống camera an ninh xử lý hình ảnh con người, quỹ đạo di chuyển và đặc trưng khuôn mặt, do đó rủi ro không chỉ nằm ở độ chính xác của mô hình mà còn ở cách dữ liệu được thu thập, lưu trữ và chia sẻ. Khi triển khai ngoài môi trường thử nghiệm, nguyên tắc tối thiểu hóa dữ liệu cần được ưu tiên: chỉ thu thập thành phần thật sự cần cho mục tiêu cảnh báo, giới hạn thời gian lưu và tránh giữ video hoặc ảnh khuôn mặt lâu hơn nhu cầu vận hành.",
    )

    add_heading(doc, "5.x.1 Dữ liệu khuôn mặt và sự đồng thuận", 3)
    add_body(
        doc,
        "Ảnh enrollment và embedding khuôn mặt là dữ liệu nhạy cảm. Người được đăng ký phải biết mục đích sử dụng, phạm vi camera, thời hạn lưu và cách yêu cầu xóa dữ liệu. Trong môi trường thử nghiệm, danh sách danh tính nên giới hạn ở người tự nguyện tham gia; khi báo cáo hoặc trình diễn, ảnh nhận dạng cần được che hoặc thay bằng dữ liệu đã được phép công bố.",
    )
    add_body(
        doc,
        "Embedding không nên được xem là dữ liệu vô danh tuyệt đối vì nó vẫn được dùng để liên kết các quan sát với một danh tính. Cơ sở dữ liệu embedding cần được tách khỏi video, áp dụng phân quyền và có cơ chế xóa đồng bộ giữa ảnh enrollment, embedding và bản ghi liên quan khi danh tính không còn được sử dụng.",
    )

    add_heading(doc, "5.x.2 Kiểm soát truy cập và bảo vệ kênh cảnh báo", 3)
    add_body(
        doc,
        "Giao diện quản trị phải phân biệt người chỉ được xem luồng camera với người có quyền thêm danh tính, thay đổi ROI, chỉnh ngưỡng hoặc xóa dữ liệu. Các thao tác ảnh hưởng đến cấu hình cảnh báo cần được ghi log gồm tài khoản, thời điểm, giá trị trước và sau thay đổi. Mật khẩu, token Telegram/Discord và thông tin camera không được ghi trực tiếp vào báo cáo, log công khai hoặc mã nguồn được chia sẻ.",
    )
    add_body(
        doc,
        "Thông báo gửi ra ngoài hệ thống có thể chứa ảnh, thời gian và vị trí camera. Do đó, nội dung cần được giới hạn ở mức đủ để xử lý sự cố; kênh gửi phải có danh sách người nhận rõ ràng và cơ chế thu hồi token khi thành viên rời nhóm. Retry phải dùng khóa chống trùng để tránh gửi nhiều cảnh báo giống nhau khi mạng chập chờn.",
    )

    add_heading(doc, "5.x.3 Giám sát của con người và giới hạn sử dụng", 3)
    add_body(
        doc,
        "Cảnh báo của hệ thống nên được xem là tín hiệu hỗ trợ ra quyết định, không phải kết luận tự động về ý định hoặc hành vi của một người. Các nhãn như 'đáng ngờ' có thể chịu ảnh hưởng từ ngưỡng, vị trí camera và đặc điểm môi trường; vì vậy người vận hành cần xem lại đoạn video liên quan trước khi thực hiện hành động có ảnh hưởng đến cá nhân.",
    )
    add_body(
        doc,
        "Hiệu năng có thể khác nhau theo khoảng cách, góc mặt, điều kiện ánh sáng và mức che khuất. Khi mở rộng dữ liệu, cần kiểm tra chênh lệch FAR/FRR giữa các nhóm điều kiện thay vì chỉ báo cáo một metric tổng. Nếu không có đủ dữ liệu để đánh giá công bằng giữa các nhóm, báo cáo phải nêu rõ giới hạn này.",
    )

    add_table(
        doc,
        "Bảng 5-x: Rủi ro triển khai và biện pháp kiểm soát",
        ["Rủi ro", "Hậu quả", "Biện pháp kiểm soát đề xuất"],
        [
            ["Truy cập trái phép", "Lộ video, ảnh hoặc embedding khuôn mặt.", "Phân quyền, xác thực, log truy cập và giới hạn thời gian lưu."],
            ["Rò rỉ token", "Kẻ khác gửi/đọc cảnh báo trên kênh tích hợp.", "Tách secret khỏi mã nguồn, luân chuyển và thu hồi token."],
            ["Giả mạo đầu vào", "Ảnh/video giả làm sai nhận diện hoặc cảnh báo.", "Giới hạn nguồn camera, kiểm tra luồng và bổ sung liveness khi cần."],
            ["Sửa cấu hình", "ROI hoặc ngưỡng bị thay đổi ngoài kiểm soát.", "Phân quyền sửa, lịch sử thay đổi và khả năng khôi phục cấu hình."],
            ["Cảnh báo sai", "Ảnh hưởng đến người bị ghi hình.", "Xác minh bởi con người và lưu bằng chứng trước khi hành động."],
        ],
        [3.0, 4.8, 7.7],
    )

    add_heading(doc, "5.x Mối đe dọa đối với tính hợp lệ", 2)
    add_body(
        doc,
        "Phân tích mối đe dọa đối với tính hợp lệ giúp xác định phạm vi mà kết quả có thể được tin cậy và suy rộng. Phần này không phủ nhận kết quả thực nghiệm; mục tiêu là tách rõ điều hệ thống đã chứng minh trong phạm vi dữ liệu hiện có với những kết luận vẫn cần thêm bằng chứng.",
    )

    add_heading(doc, "5.x.1 Tính hợp lệ nội tại", 3)
    add_body(
        doc,
        "Kết quả có thể bị ảnh hưởng bởi tải nền của hệ điều hành, trạng thái warm-up của GPU, cache, luồng gửi thông báo và thứ tự video. Nếu các yếu tố này không được kiểm soát, chênh lệch giữa hai cấu hình có thể đến từ môi trường đo thay vì thay đổi thuật toán. Biện pháp giảm thiểu là khóa cấu hình, lặp phép đo, ghi trạng thái phần cứng và sử dụng cùng một cửa sổ đo cho mọi cấu hình.",
    )

    add_heading(doc, "5.x.2 Tính hợp lệ cấu trúc", 3)
    add_body(
        doc,
        "Một metric đơn lẻ không đại diện cho toàn bộ chất lượng hệ thống. FPS cao không đảm bảo cảnh báo đúng; recall cao có thể đi kèm nhiều cảnh báo giả; FAR/FRR ở mức ảnh không phản ánh hoàn toàn độ ổn định nhãn ở mức track. Vì vậy, báo cáo kết hợp metric phát hiện, tracking, danh tính, sự kiện và độ trễ, đồng thời phân tích các ca lỗi đại diện.",
    )

    add_heading(doc, "5.x.3 Tính hợp lệ bên ngoài", 3)
    add_body(
        doc,
        "Pilot hành vi hiện có một người tham gia và pilot khuôn mặt chỉ có hai danh tính, nên chưa đại diện cho đa dạng ngoại hình, quần áo, góc quay và kiểu chuyển động trong thực tế. Các video chủ yếu phản ánh số ít địa điểm và camera; do đó chưa thể suy rộng trực tiếp sang môi trường đông người, ngoài trời phức tạp hoặc hệ thống nhiều camera hơn cấu hình đã đo.",
    )

    add_heading(doc, "5.x.4 Tính hợp lệ kết luận", 3)
    add_body(
        doc,
        "Sample size nhỏ làm các tỉ lệ thay đổi mạnh khi chỉ một mẫu được phân loại khác đi. Ví dụ, với tám ảnh frozen test khuôn mặt, một lỗi tương ứng thay đổi 0,125 trên tỉ lệ. Vì vậy, ROC, EER, FAR và FRR hiện phù hợp để mô tả pilot và lựa chọn ngưỡng sơ bộ, chưa đủ để khẳng định hiệu năng trên dân số lớn.",
    )

    add_table(
        doc,
        "Bảng 5-x: Tóm tắt mối đe dọa đối với tính hợp lệ",
        ["Loại", "Nguy cơ chính", "Hướng giảm thiểu"],
        [
            ["Nội tại", "Tải nền, warm-up, cache và cấu hình thay đổi.", "Khóa môi trường, lặp phép đo, ghi log tài nguyên."],
            ["Cấu trúc", "Metric không phản ánh đúng chất lượng đầu cuối.", "Kết hợp metric box, track, sự kiện, danh tính và độ trễ."],
            ["Bên ngoài", "Ít người, ít camera và ít điều kiện môi trường.", "Mở rộng participant, địa điểm, session và điều kiện ánh sáng."],
            ["Kết luận", "Sample size nhỏ và bước thay đổi metric lớn.", "Báo số đếm gốc, độ phân tán và giới hạn suy rộng."],
        ],
        [2.6, 6.0, 6.9],
    )

    add_heading(doc, "5.x.5 Nguyên tắc diễn giải và công bố kết quả", 3)
    add_body(
        doc,
        "Khi trình bày kết quả, báo cáo nên tách rõ ba mức: quan sát trực tiếp trên dữ liệu, giải thích kỹ thuật về nguyên nhân và phạm vi có thể suy rộng. Mỗi nhận định cần gắn với tập dữ liệu, số lần chạy, cấu hình và metric tương ứng. Cụm từ ‘không quan sát thấy lỗi’ chỉ mô tả những mẫu đã kiểm tra, không đồng nghĩa với việc lỗi không thể xảy ra ở điều kiện khác.",
    )
    add_body(
        doc,
        "Các lần chạy thất bại, bị hủy hoặc có tài nguyên bất thường vẫn cần được ghi trong nhật ký thí nghiệm và nêu lý do loại khỏi phép tổng hợp. Không nên chỉ chọn lần chạy có kết quả tốt nhất. Nếu một kết luận thay đổi khi loại bỏ một mẫu hoặc một lần chạy, báo cáo phải nêu rõ độ nhạy này và xem kết quả là bằng chứng sơ bộ cần được kiểm chứng thêm.",
    )

    page_break(doc)
    add_heading(doc, "PHỤ LỤC A: MẪU GHI NHẬN THÍ NGHIỆM VÀ CA LỖI", 1)
    add_body(
        doc,
        "Phụ lục này cung cấp biểu mẫu thống nhất để lưu thông tin của mỗi phép đo và ca lỗi. Khi đưa vào báo cáo chính, có thể giữ một mẫu hoàn chỉnh đại diện và chuyển phần log chi tiết sang tệp dữ liệu đi kèm.",
    )

    add_heading(doc, "A.1 Phiếu ghi nhận một lần chạy", 2)
    add_table(
        doc,
        "Bảng A-1: Phiếu ghi nhận cấu hình và kết quả một lần chạy",
        ["Trường", "Giá trị cần điền"],
        [
            ["Mã lần chạy", "[[BỔ SUNG]]"],
            ["Ngày giờ bắt đầu/kết thúc", "[[BỔ SUNG]]"],
            ["Mã cấu hình ablation", "[[BỔ SUNG]]"],
            ["Commit và checkpoint", "[[BỔ SUNG]]"],
            ["Dữ liệu đầu vào", "[[BỔ SUNG]]"],
            ["Phần cứng/phần mềm", "[[BỔ SUNG]]"],
            ["Ngưỡng và file cấu hình", "[[BỔ SUNG]]"],
            ["Warm-up và cửa sổ đo", "[[BỔ SUNG]]"],
            ["IDF1/ID switch/fragmentation", "[[BỔ SUNG]]"],
            ["FPS/p50/p95/p99/VRAM", "[[BỔ SUNG]]"],
            ["Sự cố bất thường", "[[BỔ SUNG hoặc ghi 'Không']]"],
        ],
        [5.0, 10.5],
        placeholder_columns=(1,),
    )

    add_heading(doc, "A.2 Phiếu phân tích một ca lỗi", 2)
    add_table(
        doc,
        "Bảng A-2: Phiếu phân tích ca lỗi theo chuỗi xử lý",
        ["Trường", "Nội dung"],
        [
            ["Mã ca lỗi và video", "[[BỔ SUNG]]"],
            ["Khung hình/thời gian xảy ra", "[[BỔ SUNG]]"],
            ["Ground truth", "[[BỔ SUNG]]"],
            ["Kết quả hệ thống", "[[BỔ SUNG]]"],
            ["Tầng đầu tiên tạo sai lệch", "[[BỔ SUNG]]"],
            ["Confidence/IoU/cosine score", "[[BỔ SUNG]]"],
            ["Ảnh trước - trong - sau lỗi", "[[CHÈN HÌNH]]"],
            ["Nguyên nhân giả thuyết", "[[BỔ SUNG]]"],
            ["Biện pháp khắc phục", "[[BỔ SUNG]]"],
            ["Ảnh hưởng phụ cần kiểm tra", "[[BỔ SUNG]]"],
        ],
        [5.0, 10.5],
        placeholder_columns=(1,),
    )

    final_note = doc.add_paragraph(style="Normal")
    final_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    final_note.paragraph_format.first_line_indent = Cm(0)
    final_note.paragraph_format.line_spacing = 1.15
    final_note.paragraph_format.space_after = Pt(10)
    shade_paragraph(final_note, "E2F0D9")
    set_run_font(final_note.add_run("Kiểm tra trước khi chèn: "), 12, bold=True, italic=True)
    set_run_font(
        final_note.add_run(
            "thay toàn bộ ký hiệu 4.x/5.x bằng số mục thực tế; thay mọi ô tô vàng bằng dữ liệu thật; "
            "cập nhật Caption, mục lục, danh mục bảng và các tham chiếu chéo sau khi dán vào báo cáo chính."
        ),
        12,
        italic=True,
    )

    doc.save(OUTPUT)

    check = Document(OUTPUT)
    text_parts = [p.text for p in check.paragraphs]
    text_parts.extend(
        p.text
        for table in check.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )
    full_text = "\n".join(text_parts)
    assert "TRƯỜNG ĐẠI HỌC" not in full_text
    assert "Đánh giá đóng góp của từng thành phần" in full_text
    assert "Quyền riêng tư, bảo mật và đạo đức triển khai" in full_text
    assert len(check.tables) >= 10
    assert full_text.count("[[BỔ SUNG") >= 10
    assert abs(check.sections[0].left_margin.cm - 3.5) < 0.05
    print(f"Created {OUTPUT}")
    print(f"Paragraphs={len(check.paragraphs)} Tables={len(check.tables)} Placeholders={full_text.count('[[BỔ SUNG')}")


if __name__ == "__main__":
    build()
